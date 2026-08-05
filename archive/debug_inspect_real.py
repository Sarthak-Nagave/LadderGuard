"""Inspect the real Ladder_Chronology.xlsx to understand its structure."""
import re

import openpyxl

paths = [
    r"C:\Users\SM 464\Desktop\Operational Package\7. Chronology\Master\Initial\Ladder_Chronology.xlsx",
    r"C:\Users\SM 464\Desktop\Operational Package\7. Chronology\Slave\Final\Ladder_Chronology.xlsx",
    r"C:\Users\SM 464\Desktop\Operational Package\7. Chronology\UUT\QC\Ladder_Chronology.xlsx",
]

def normalize(val):
    if val is None:
        return ""
    text = str(val).replace("\n", " ").replace("\r", " ").replace("\t", " ")
    text = text.replace("\u2013", "-").replace("\u2014", "-")
    text = text.lower()
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"\s*-\s*", "-", text)
    return text.strip()

for path in paths:
    try:
        wb = openpyxl.load_workbook(path)
    except Exception as e:
        print(f"SKIP {path}: {e}")
        continue
    ws = wb.active
    print(f"\n{'=' * 100}")
    print(f"FILE: {path}")
    print(f"Sheet: {ws.title}, max_row={ws.max_row}, max_col={ws.max_column}")

    print(f"\nMerged ranges ({len(list(ws.merged_cells.ranges))}):")
    for mr in ws.merged_cells.ranges:
        col_span = mr.max_col - mr.min_col + 1
        row_span = mr.max_row - mr.min_row + 1
        master_val = ws.cell(row=mr.min_row, column=mr.min_col).value
        print(f"  {mr!s:20s}  rows={mr.min_row}-{mr.max_row} cols={mr.min_col}-{mr.max_col} (span {col_span}x{row_span})  master_val={repr(master_val)[:50]}")

    # Find header row
    header_row = None
    for r in range(1, 51):
        for c in range(1, ws.max_column + 1):
            if normalize(ws.cell(row=r, column=c).value) == "serial no.":
                header_row = r
                break
        if header_row:
            break

    print(f"\nHeader row: {header_row}")
    if not header_row:
        continue

    # Build header map
    header_map = {}
    for c in range(1, ws.max_column + 1):
        val = ws.cell(row=header_row, column=c).value
        norm = normalize(val)
        if norm:
            header_map[norm] = c
        raw_repr = repr(val)
        if len(raw_repr) > 50:
            raw_repr = raw_repr[:50] + "..."
        print(f"  Col {c:2d}: norm={norm:45s} raw={raw_repr}")

    # Print data rows
    first_data = header_row + 1
    print(f"\nData rows ({first_data} to {min(ws.max_row, header_row + 10)}):")
    for r in range(first_data, min(ws.max_row + 1, header_row + 10)):
        # Check if this row is in a merged range
        in_merged = False
        for mr in ws.merged_cells.ranges:
            if mr.min_row <= r <= mr.max_row and (mr.max_col - mr.min_col) >= 2:
                in_merged = True
                break

        print(f"\n  Row {r} {'[MERGED/INSTRUCTION]' if in_merged else ''}:")
        for header, col in sorted(header_map.items(), key=lambda x: x[1]):
            cell = ws.cell(row=r, column=col)
            val = cell.value
            # Check if this cell is in a merged range
            cell_in_merge = False
            for mr in ws.merged_cells.ranges:
                if mr.min_row <= r <= mr.max_row and mr.min_col <= col <= mr.max_col:
                    cell_in_merge = True
                    break
            merge_tag = " [M]" if cell_in_merge else ""
            val_repr = repr(val) if val is not None else "<<EMPTY>>"
            if len(val_repr) > 60:
                val_repr = val_repr[:60] + "..."
            print(f"    Col {col:2d} {header:42s} = {val_repr}{merge_tag}")

    wb.close()

print("\n\nDone.")
