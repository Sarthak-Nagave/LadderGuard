"""
Debug script: Trace the exact runtime behavior of ChronologyExcelWriter
step-by-step when updating an EXISTING workbook with a new release.

This script does NOT modify excel_writer.py.
It replicates the exact write() code path, pausing after every step
to inspect the worksheet state.

Goal: Determine EXACTLY when and WHERE previous release values disappear.
"""

from __future__ import annotations

import sys
import tempfile
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

# ──────────────────────────────────────────────────────────────────────
# HEADERS (must match excel_writer.py exactly)
# ──────────────────────────────────────────────────────────────────────
HEADERS = [
    "Serial No.",
    "Source Code Path",
    "Bin File Name",
    "CRC",
    "Ladder Version No.",
    "Reason for Upgrade",
    "Testing Stage",
    "PLC Model",
    "Selpro Version & Path",
    "Bootloader Version",
    "Release Date",
    "Released By",
    "Ladder Release-To Production",
    "Operator Procedure Modification",
    "Automation Set Up Modification",
    "Tested By",
]

# A "real" template has headers on row 1.  We will also test with headers
# on a later row (row 15, mimicking the real-world template that has a
# title/instructions block above the data table).

SEPARATOR = "=" * 90


def build_header_map(sheet: Worksheet, header_row: int) -> dict[str, int]:
    """Identical logic to ChronologyExcelWriter._build_header_map"""
    import re
    header_map: dict[str, int] = {}
    for col_idx in range(1, sheet.max_column + 1):
        value = sheet.cell(row=header_row, column=col_idx).value
        if value is None:
            continue
        text = str(value)
        text = text.replace("\n", " ").replace("\r", " ").replace("\t", " ")
        text = text.replace("–", "-").replace("—", "-")
        text = text.lower()
        text = re.sub(r"\s+", " ", text)
        text = re.sub(r"\s*-\s*", "-", text)
        text = text.strip()
        if text:
            header_map[text] = col_idx
    return header_map


def snapshot_row(sheet: Worksheet, row: int, header_map: dict[str, int]) -> dict[str, str | None]:
    """Read every mapped cell in a row and return {header: value}."""
    result = {}
    for header, col in header_map.items():
        cell = sheet.cell(row=row, column=col)
        result[header] = cell.value
    return result


def snapshot_row_styles(sheet: Worksheet, row: int, header_map: dict[str, int]) -> dict[str, dict]:
    """Read style attributes for every cell in a row."""
    result = {}
    for header, col in header_map.items():
        cell = sheet.cell(row=row, column=col)
        result[header] = {
            "font_bold": cell.font.bold if cell.font else None,
            "font_name": cell.font.name if cell.font else None,
            "font_size": cell.font.size if cell.font else None,
            "fill_color": cell.fill.fgColor.rgb if (cell.fill and cell.fill.fgColor) else None,
            "border_left": cell.border.left.style if (cell.border and cell.border.left) else None,
            "border_right": cell.border.right.style if (cell.border and cell.border.right) else None,
            "border_top": cell.border.top.style if (cell.border and cell.border.top) else None,
            "border_bottom": cell.border.bottom.style if (cell.border and cell.border.bottom) else None,
            "alignment_h": cell.alignment.horizontal if cell.alignment else None,
            "alignment_v": cell.alignment.vertical if cell.alignment else None,
            "number_format": cell.number_format,
        }
    return result


def print_merged_ranges(sheet: Worksheet, label: str):
    """Print all merged cell ranges."""
    ranges = list(sheet.merged_cells.ranges)
    print(f"\n  Merged ranges ({label}): {len(ranges)}")
    for mr in ranges:
        print(f"    {mr}")


def print_row_snapshot(label: str, row_num: int, snap: dict):
    """Pretty print a row snapshot."""
    print(f"\n  {label} (Row {row_num}):")
    for header, value in snap.items():
        display = repr(value) if value is not None else "<<EMPTY/None>>"
        print(f"    {header:>42s} = {display}")


def print_row_styles(label: str, row_num: int, styles: dict):
    """Pretty print style info for a row."""
    print(f"\n  Styles {label} (Row {row_num}):")
    for header, s in styles.items():
        parts = []
        if s["font_bold"]:
            parts.append("BOLD")
        if s["border_left"] or s["border_right"] or s["border_top"] or s["border_bottom"]:
            parts.append(f"borders=[L:{s['border_left']} R:{s['border_right']} T:{s['border_top']} B:{s['border_bottom']}]")
        if s["fill_color"] and s["fill_color"] != "00000000":
            parts.append(f"fill={s['fill_color']}")
        if s["alignment_h"]:
            parts.append(f"align_h={s['alignment_h']}")
        if parts:
            print(f"    {header:>42s} : {', '.join(parts)}")


def inspect_state(sheet: Worksheet, header_row: int, header_map: dict[str, int], title: str, max_data_rows: int = 5):
    """Full inspection of the worksheet state."""
    print(f"\n{SEPARATOR}")
    print(f"  STATE: {title}")
    print(SEPARATOR)
    print(f"  sheet.max_row = {sheet.max_row}")
    print(f"  sheet.max_column = {sheet.max_column}")
    print(f"  header_row = {header_row}")

    print_merged_ranges(sheet, title)

    for offset in range(max_data_rows):
        row = header_row + 1 + offset
        if row > sheet.max_row:
            break
        snap = snapshot_row(sheet, row, header_map)
        has_any = any(v is not None and str(v).strip() != "" for v in snap.values())
        if not has_any and offset > 0:
            print(f"\n  Row {row}: <<ALL EMPTY>> — stopping scan.")
            break
        print_row_snapshot("DATA ROW", row, snap)
        print_row_styles("DATA ROW", row, snapshot_row_styles(sheet, row, header_map))

    print(SEPARATOR)


# ──────────────────────────────────────────────────────────────────────
# CREATE A REALISTIC TEMPLATE (headers at row 1, simple case)
# ──────────────────────────────────────────────────────────────────────
def create_template(path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    header_font = Font(bold=True, size=11, name="Calibri")
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    for idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wb.save(path)


# ──────────────────────────────────────────────────────────────────────
# POPULATE V1.00 RELEASE (simulate the FIRST generation)
# ──────────────────────────────────────────────────────────────────────
def populate_v100(path: Path):
    """Simulate a completed V1.00 release row — exactly as the real writer does."""
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    hmap = build_header_map(ws, 1)
    row = 2  # first data row

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    data_font = Font(bold=False, size=10, name="Calibri")

    values = {
        "serial no.": 1,
        "source code path": r"C:\Project\2. Bin File\Master\Initial\FW_MIBRX-4M_BL20_V1.00.bin",
        "bin file name": "FW_MIBRX-4M_BL20_V1.00.bin",
        "crc": "A3F91C7E",
        "ladder version no.": "V1.00",
        "reason for upgrade": "Initial Release",
        "testing stage": "Master Initial",
        "plc model": "MIBRX-4M",
        "selpro version & path": "SP3.5 C:/Selpro",
        "bootloader version": "BL20",
        "release date": "01/07/2026",
        "released by": "Engineer A",
        "ladder release-to production": "Yes",
        "operator procedure modification": "No",
        "automation set up modification": "No",
        "tested by": "QA Team A",
    }

    for header, col_idx in hmap.items():
        cell = ws.cell(row=row, column=col_idx)
        if header in values:
            cell.value = values[header]
        cell.font = data_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb.save(path)


# ──────────────────────────────────────────────────────────────────────
# CREATE TEMPLATE WITH MERGED CELLS (more realistic scenario)
# ──────────────────────────────────────────────────────────────────────
def create_template_with_merges(path: Path):
    """
    Creates a template that is closer to the real-world Ladder Chronology:
    - Rows 1-14: Title block with some merged cells
    - Row 15: Header row
    - Row 16: Instruction/merged row (spanning many columns)
    - Row 17+: Data rows
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    header_font = Font(bold=True, size=11, name="Calibri")
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    # Title in row 1, merged across columns
    ws.cell(row=1, column=1).value = "LADDER CHRONOLOGY"
    ws.merge_cells("A1:P1")

    # Headers at row 15
    for idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=15, column=idx)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Instruction row at row 16 — merged across many columns (structural merge)
    ws.cell(row=16, column=1).value = "Fill in the fields below for each release"
    ws.merge_cells("A16:P16")

    wb.save(path)


def populate_v100_with_merges(path: Path, header_row: int = 15, data_row: int = 17):
    """Populate V1.00 in the merged-cell template."""
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    hmap = build_header_map(ws, header_row)

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    data_font = Font(bold=False, size=10, name="Calibri")

    values = {
        "serial no.": 1,
        "source code path": r"C:\Project\2. Bin File\Master\Initial\FW_MIBRX-4M_BL20_V1.00.bin",
        "bin file name": "FW_MIBRX-4M_BL20_V1.00.bin",
        "crc": "A3F91C7E",
        "ladder version no.": "V1.00",
        "reason for upgrade": "Initial Release",
        "testing stage": "Master Initial",
        "plc model": "MIBRX-4M",
        "selpro version & path": "SP3.5 C:/Selpro",
        "bootloader version": "BL20",
        "release date": "01/07/2026",
        "released by": "Engineer A",
        "ladder release-to production": "Yes",
        "operator procedure modification": "No",
        "automation set up modification": "No",
        "tested by": "QA Team A",
    }

    for header, col_idx in hmap.items():
        cell = ws.cell(row=data_row, column=col_idx)
        if header in values:
            cell.value = values[header]
        cell.font = data_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb.save(path)


# ──────────────────────────────────────────────────────────────────────
# IMPORT AND USE THE ACTUAL ChronologyExcelWriter METHODS
# ──────────────────────────────────────────────────────────────────────
# We import the actual class and call its internal methods one-by-one,
# inspecting the worksheet state after each call.
sys.path.insert(0, str(Path(__file__).resolve().parent))
from services.chronology_generator.excel_writer import ChronologyExcelWriter
from services.chronology_generator.models import ChronologyEntry


def run_trace(template_path: Path, output_path: Path, label: str):
    """
    Trace the EXACT execution flow of ChronologyExcelWriter.write()
    step-by-step, inspecting the worksheet after every operation.
    """
    print(f"\n\n{'#' * 90}")
    print(f"# TRACE: {label}")
    print(f"# template = {template_path}")
    print(f"# output   = {output_path}")
    print(f"{'#' * 90}")

    # ── Step 0: Create the writer (same as production code) ──
    writer = ChronologyExcelWriter(template_path)

    # ── Step 1: Load workbook ──
    print("\n>>> STEP 1: _load_workbook()")
    wb = writer._load_workbook()
    ws = wb.active
    assert ws is not None

    # ── Step 2: Find header row and build header map ──
    print(">>> STEP 2: _find_header_row() + _build_header_map()")
    header_row_idx = writer._find_header_row(ws)
    header_map = writer._build_header_map(ws, header_row_idx)
    first_data_row = writer._first_data_row(header_row_idx)
    print(f"  header_row_idx   = {header_row_idx}")
    print(f"  first_data_row   = {first_data_row}")
    print(f"  header_map keys  = {sorted(header_map.keys())}")

    # ── Step 3: Skip instruction rows ──
    history_row = first_data_row
    while writer._is_instruction_row(ws, history_row, header_map):
        print(f"  Skipping instruction row {history_row}")
        history_row += 1
    print(f"  history_row (first real data row) = {history_row}")

    # ── Step 4: Detect existing workbook / history ──
    using_existing = writer._is_existing_workbook(output_path)
    has_history = writer._has_history_row_data(ws, history_row, header_map)
    print("\n>>> STEP 4: Detect existing workbook")
    print(f"  using_existing_workbook  = {using_existing}")
    print(f"  has_existing_history     = {has_history}")

    inspect_state(ws, header_row_idx, header_map, "BEFORE ANY MODIFICATION", max_data_rows=6)

    # ── Record pre-insertion state of previous rows ──
    pre_insertion_snapshots = {}
    for offset in range(10):
        r = history_row + offset
        if r > ws.max_row:
            break
        snap = snapshot_row(ws, r, header_map)
        if any(v is not None and str(v).strip() for v in snap.values()):
            pre_insertion_snapshots[r] = snap

    print(f"\n  Pre-insertion populated rows: {sorted(pre_insertion_snapshots.keys())}")

    # ── Step 5: insert_rows() ──
    if using_existing and has_history:
        print(f"\n>>> STEP 5: sheet.insert_rows({history_row}, amount=1)")
        ws.insert_rows(history_row, amount=1)
        inspect_state(ws, header_row_idx, header_map, "AFTER insert_rows()", max_data_rows=6)

        # ── Check: did insert_rows() preserve old row data? ──
        print("\n  ── VERIFICATION: Did insert_rows() preserve old row values? ──")
        for old_row, old_snap in pre_insertion_snapshots.items():
            new_row = old_row + 1  # shifted down by 1
            new_snap = snapshot_row(ws, new_row, header_map)
            for header in old_snap:
                old_val = old_snap[header]
                new_val = new_snap.get(header)
                if old_val != new_val:
                    print(f"    !! MISMATCH at row {old_row}->{new_row}, col '{header}': was {old_val!r}, now {new_val!r}")
            all_match = all(old_snap[h] == new_snap.get(h) for h in old_snap)
            if all_match:
                print(f"    Row {old_row} -> {new_row}: ALL VALUES PRESERVED ✓")
            else:
                print(f"    Row {old_row} -> {new_row}: SOME VALUES LOST ✗")

        # ── Step 6: _copy_row_format() ──
        print(f"\n>>> STEP 6: _copy_row_format(sheet, source_row={history_row + 1}, target_row={history_row})")
        writer._copy_row_format(ws, history_row + 1, history_row)
        inspect_state(ws, header_row_idx, header_map, "AFTER _copy_row_format()", max_data_rows=6)

        # ── Check: did _copy_row_format() corrupt old row data? ──
        print("\n  ── VERIFICATION: Did _copy_row_format() preserve old row values? ──")
        for old_row, old_snap in pre_insertion_snapshots.items():
            new_row = old_row + 1
            new_snap = snapshot_row(ws, new_row, header_map)
            for header in old_snap:
                old_val = old_snap[header]
                new_val = new_snap.get(header)
                if old_val != new_val:
                    print(f"    !! MISMATCH at row {old_row}->{new_row}, col '{header}': was {old_val!r}, now {new_val!r}")
            all_match = all(old_snap[h] == new_snap.get(h) for h in old_snap)
            if all_match:
                print(f"    Row {old_row} -> {new_row}: ALL VALUES PRESERVED ✓")
            else:
                print(f"    Row {old_row} -> {new_row}: SOME VALUES LOST ✗")

        # Also check: did _copy_row_format() write any VALUES into the new row?
        new_row_snap = snapshot_row(ws, history_row, header_map)
        print(f"\n  ── New row {history_row} after _copy_row_format(): ──")
        for header, val in new_row_snap.items():
            if val is not None and str(val).strip():
                print(f"    !! UNEXPECTED VALUE in new row: '{header}' = {val!r}")
    else:
        print("\n>>> STEP 5-6: SKIPPED (no existing history or not existing workbook)")

    # ── Step 7: _write_entry() ──
    print(f"\n>>> STEP 7: _write_entry() into row {history_row}")
    new_entry = ChronologyEntry(
        source_code_path=Path(r"C:\Project\2. Bin File\Master\Initial\FW_MIBRX-4M_BL20_V1.01.bin"),
        bin_file_path=Path(r"C:\Project\2. Bin File\Master\Initial\FW_MIBRX-4M_BL20_V1.01.bin"),
        sdoc_file_path=None,
        bin_file_name="FW_MIBRX-4M_BL20_V1.01.bin",
        sdoc_file_name=None,
        version="V1.01",
        plc_model="MIBRX-4M",
        selpro_version="SP4.0 C:/Selpro",
        bootloader_version="BL20",
        crc="B4C82D9F",
        testing_stage="Master Initial",
        release_date="",
        reason_for_upgrade="Bug Fix Release",
        released_by="Engineer B",
        tested_by="QA Team B",
        ladder_release_to_production="Yes",
        operator_procedure_modification="No",
        automation_setup_modification="No",
    )
    release_date = datetime.now().strftime("%d/%m/%Y")
    writer._write_entry(ws, history_row, new_entry, header_map, serial_no=1, release_date=release_date)
    inspect_state(ws, header_row_idx, header_map, "AFTER _write_entry()", max_data_rows=6)

    # ── Check: did _write_entry() corrupt old row data? ──
    if pre_insertion_snapshots:
        print("\n  ── VERIFICATION: Did _write_entry() preserve old row values? ──")
        for old_row, old_snap in pre_insertion_snapshots.items():
            new_row = old_row + 1
            new_snap = snapshot_row(ws, new_row, header_map)
            for header in old_snap:
                old_val = old_snap[header]
                new_val = new_snap.get(header)
                if old_val != new_val:
                    print(f"    !! MISMATCH at row {old_row}->{new_row}, col '{header}': was {old_val!r}, now {new_val!r}")
            all_match = all(old_snap[h] == new_snap.get(h) for h in old_snap)
            if all_match:
                print(f"    Row {old_row} -> {new_row}: ALL VALUES PRESERVED ✓")
            else:
                print(f"    Row {old_row} -> {new_row}: SOME VALUES LOST ✗")

    # ── Step 8: _renumber_serials() ──
    print(f"\n>>> STEP 8: _renumber_serials() starting at row {history_row}")
    writer._renumber_serials(ws, history_row, header_map)
    inspect_state(ws, header_row_idx, header_map, "AFTER _renumber_serials()", max_data_rows=6)

    # ── Check: did _renumber_serials() corrupt old row data? ──
    if pre_insertion_snapshots:
        print("\n  ── VERIFICATION: Did _renumber_serials() preserve old row values? ──")
        for old_row, old_snap in pre_insertion_snapshots.items():
            new_row = old_row + 1
            new_snap = snapshot_row(ws, new_row, header_map)
            for header in old_snap:
                if header == "serial no.":
                    continue  # serial is expected to change
                old_val = old_snap[header]
                new_val = new_snap.get(header)
                if old_val != new_val:
                    print(f"    !! MISMATCH at row {old_row}->{new_row}, col '{header}': was {old_val!r}, now {new_val!r}")
            non_serial = {h: old_snap[h] for h in old_snap if h != "serial no."}
            non_serial_new = {h: new_snap.get(h) for h in old_snap if h != "serial no."}
            if non_serial == non_serial_new:
                print(f"    Row {old_row} -> {new_row}: ALL NON-SERIAL VALUES PRESERVED ✓")
            else:
                print(f"    Row {old_row} -> {new_row}: SOME VALUES LOST ✗")

    # ── Step 9: Save ──
    print(f"\n>>> STEP 9: workbook.save({output_path})")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    # ── Step 10: Reload and verify ──
    print("\n>>> STEP 10: Reload saved workbook and verify")
    wb2 = openpyxl.load_workbook(output_path)
    ws2 = wb2.active
    header_row_idx2 = header_row_idx  # same header row
    header_map2 = build_header_map(ws2, header_row_idx2)
    inspect_state(ws2, header_row_idx2, header_map2, "AFTER SAVE + RELOAD", max_data_rows=6)

    if pre_insertion_snapshots:
        print("\n  ── VERIFICATION: Did save+reload preserve old row values? ──")
        for old_row, old_snap in pre_insertion_snapshots.items():
            new_row = old_row + 1
            new_snap = snapshot_row(ws2, new_row, header_map2)
            for header in old_snap:
                if header == "serial no.":
                    continue
                old_val = old_snap[header]
                new_val = new_snap.get(header)
                if old_val != new_val:
                    print(f"    !! MISMATCH at row {old_row}->{new_row}, col '{header}': was {old_val!r}, now {new_val!r}")
            non_serial = {h: old_snap[h] for h in old_snap if h != "serial no."}
            non_serial_new = {h: new_snap.get(h) for h in old_snap if h != "serial no."}
            if non_serial == non_serial_new:
                print(f"    Row {old_row} -> {new_row}: ALL NON-SERIAL VALUES PRESERVED ✓")
            else:
                print(f"    Row {old_row} -> {new_row}: SOME VALUES LOST ✗")

    return output_path


def run_multi_release_trace(template_path: Path, output_path: Path, label_prefix: str, header_row: int = 1):
    """
    Run 4 successive releases: V1.00 -> V1.01 -> V2.00 -> V3.00
    to verify cumulative history preservation.
    """
    releases = [
        ("V1.01", "B4C82D9F", "Bug Fix Release", "Engineer B", "QA Team B"),
        ("V2.00", "C5D93E0A", "Major Upgrade", "Engineer C", "QA Team C"),
        ("V3.00", "D6EA4F1B", "Feature Addition", "Engineer D", "QA Team D"),
    ]

    # First release is already populated in the template as V1.00.
    # Successive releases use the output as both template and output.
    current_template = template_path

    for i, (version, crc, reason, released_by, tested_by) in enumerate(releases):
        rel_label = f"{label_prefix} — Release #{i + 2} ({version})"
        writer = ChronologyExcelWriter(current_template)

        wb = writer._load_workbook()
        ws = wb.active
        hri = writer._find_header_row(ws)
        hmap = writer._build_header_map(ws, hri)
        fdr = writer._first_data_row(hri)

        history_row = fdr
        while writer._is_instruction_row(ws, history_row, hmap):
            history_row += 1

        using_existing = writer._is_existing_workbook(output_path)
        has_history = writer._has_history_row_data(ws, history_row, hmap)

        print(f"\n\n{'#' * 90}")
        print(f"# MULTI-RELEASE TRACE: {rel_label}")
        print(f"# using_existing={using_existing}, has_history={has_history}")
        print(f"{'#' * 90}")

        # Snapshot ALL existing data rows before modification
        pre_snapshots = {}
        for offset in range(20):
            r = history_row + offset
            if r > ws.max_row:
                break
            if writer._is_instruction_row(ws, r, hmap):
                continue
            snap = snapshot_row(ws, r, hmap)
            if any(v is not None and str(v).strip() for v in snap.values()):
                pre_snapshots[r] = snap
        print(f"  Pre-existing data rows: {sorted(pre_snapshots.keys())}")

        inspect_state(ws, hri, hmap, f"BEFORE {version}", max_data_rows=8)

        if using_existing and has_history:
            ws.insert_rows(history_row, amount=1)
            writer._copy_row_format(ws, history_row + 1, history_row)

        entry = ChronologyEntry(
            source_code_path=Path(rf"C:\Project\2. Bin File\Master\Initial\FW_MIBRX-4M_BL20_{version}.bin"),
            bin_file_path=Path(rf"C:\Project\2. Bin File\Master\Initial\FW_MIBRX-4M_BL20_{version}.bin"),
            sdoc_file_path=None,
            bin_file_name=f"FW_MIBRX-4M_BL20_{version}.bin",
            sdoc_file_name=None,
            version=version,
            plc_model="MIBRX-4M",
            selpro_version=f"SP{i + 4}.0 C:/Selpro",
            bootloader_version="BL20",
            crc=crc,
            testing_stage="Master Initial",
            release_date="",
            reason_for_upgrade=reason,
            released_by=released_by,
            tested_by=tested_by,
            ladder_release_to_production="Yes",
            operator_procedure_modification="No",
            automation_setup_modification="No",
        )
        rd = datetime.now().strftime("%d/%m/%Y")
        writer._write_entry(ws, history_row, entry, hmap, serial_no=1, release_date=rd)
        writer._renumber_serials(ws, history_row, hmap)

        inspect_state(ws, hri, hmap, f"AFTER {version} COMPLETE", max_data_rows=8)

        # Verify all pre-existing rows
        print(f"\n  ── FINAL VERIFICATION for {version}: ──")
        corruption_found = False
        for old_row, old_snap in pre_snapshots.items():
            new_row = old_row + 1  # shifted down
            new_snap = snapshot_row(ws, new_row, hmap)
            for header in old_snap:
                if header == "serial no.":
                    continue
                old_val = old_snap[header]
                new_val = new_snap.get(header)
                if old_val != new_val:
                    print(f"    !! CORRUPTION at row {old_row}->{new_row}, col '{header}': was {old_val!r}, now {new_val!r}")
                    corruption_found = True
            non_serial = {h: old_snap[h] for h in old_snap if h != "serial no."}
            non_serial_new = {h: new_snap.get(h) for h in old_snap if h != "serial no."}
            if non_serial == non_serial_new:
                print(f"    Row {old_row} -> {new_row}: ALL NON-SERIAL VALUES PRESERVED ✓")
            else:
                print(f"    Row {old_row} -> {new_row}: SOME VALUES LOST ✗")

        if corruption_found:
            print(f"\n  >>>>>> CORRUPTION DETECTED DURING {version} GENERATION <<<<<<")

        wb.save(output_path)

        # From now on, use output as template
        current_template = output_path

    # Final reload check
    print(f"\n\n{'=' * 90}")
    print("FINAL RELOAD CHECK AFTER ALL 4 RELEASES")
    print(f"{'=' * 90}")
    wb_final = openpyxl.load_workbook(output_path)
    ws_final = wb_final.active
    hri_f = header_row  # known header row
    # Rebuild header map from final workbook
    hmap_f = build_header_map(ws_final, hri_f)
    inspect_state(ws_final, hri_f, hmap_f, "FINAL STATE (4 releases)", max_data_rows=10)


# ──────────────────────────────────────────────────────────────────────
# MAIN
# ──────────────────────────────────────────────────────────────────────
def main():
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)

        # ════════════════════════════════════════════════════════════════
        # SCENARIO A: Simple template (headers at row 1, no merged cells)
        # ════════════════════════════════════════════════════════════════
        print("\n" + "█" * 90)
        print("█  SCENARIO A: Simple template — headers at row 1, NO merged cells")
        print("█" * 90)

        template_a = tmp / "scenario_a" / "template.xlsx"
        template_a.parent.mkdir(parents=True, exist_ok=True)
        create_template(template_a)

        output_a = tmp / "scenario_a" / "output.xlsx"
        output_a.parent.mkdir(parents=True, exist_ok=True)

        # Populate V1.00 into the output (simulate first generation)
        import shutil
        shutil.copy2(template_a, output_a)
        populate_v100(output_a)

        # Now trace V1.01 generation using output as both template and output
        run_trace(output_a, output_a, "Scenario A: V1.00 -> V1.01 (simple)")

        # ════════════════════════════════════════════════════════════════
        # SCENARIO B: Template with merged cells (closer to real-world)
        # ════════════════════════════════════════════════════════════════
        print("\n\n" + "█" * 90)
        print("█  SCENARIO B: Merged-cell template — headers at row 15, instruction row at 16")
        print("█" * 90)

        template_b = tmp / "scenario_b" / "template.xlsx"
        template_b.parent.mkdir(parents=True, exist_ok=True)
        create_template_with_merges(template_b)

        output_b = tmp / "scenario_b" / "output.xlsx"
        output_b.parent.mkdir(parents=True, exist_ok=True)

        shutil.copy2(template_b, output_b)
        populate_v100_with_merges(output_b, header_row=15, data_row=17)

        # Trace V1.01 on merged-cell template
        run_trace(output_b, output_b, "Scenario B: V1.00 -> V1.01 (merged cells)")

        # ════════════════════════════════════════════════════════════════
        # SCENARIO C: Multi-release trace (4 successive releases)
        # ════════════════════════════════════════════════════════════════
        print("\n\n" + "█" * 90)
        print("█  SCENARIO C: Multi-release — V1.00 -> V1.01 -> V2.00 -> V3.00 (simple template)")
        print("█" * 90)

        template_c = tmp / "scenario_c" / "template.xlsx"
        template_c.parent.mkdir(parents=True, exist_ok=True)
        create_template(template_c)

        output_c = tmp / "scenario_c" / "output.xlsx"
        output_c.parent.mkdir(parents=True, exist_ok=True)

        shutil.copy2(template_c, output_c)
        populate_v100(output_c)

        run_multi_release_trace(output_c, output_c, "Scenario C (simple)", header_row=1)

        # ════════════════════════════════════════════════════════════════
        # SCENARIO D: Multi-release with merged cells
        # ════════════════════════════════════════════════════════════════
        print("\n\n" + "█" * 90)
        print("█  SCENARIO D: Multi-release — V1.00 -> V1.01 -> V2.00 -> V3.00 (merged template)")
        print("█" * 90)

        template_d = tmp / "scenario_d" / "template.xlsx"
        template_d.parent.mkdir(parents=True, exist_ok=True)
        create_template_with_merges(template_d)

        output_d = tmp / "scenario_d" / "output.xlsx"
        output_d.parent.mkdir(parents=True, exist_ok=True)

        shutil.copy2(template_d, output_d)
        populate_v100_with_merges(output_d, header_row=15, data_row=17)

        run_multi_release_trace(output_d, output_d, "Scenario D (merged)", header_row=15)

    print("\n\n" + "█" * 90)
    print("█  ALL SCENARIOS COMPLETE")
    print("█" * 90)


if __name__ == "__main__":
    main()
