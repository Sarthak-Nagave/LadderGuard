import openpyxl

wb = openpyxl.Workbook()
ws = wb.active
ws.cell(row=15, column=2, value="Data")
ws.cell(row=16, column=2, value="Instruction")
ws.merge_cells('B16:P16')
print("Before:")
print(list(ws.merged_cells.ranges))
ws.insert_rows(14, 1)
print("After inserting at 14:")
print(list(ws.merged_cells.ranges))
