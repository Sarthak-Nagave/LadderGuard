"""
Debug script that exactly replicates the real Ladder_Chronology.xlsx structure:
- Header at row 14
- Merged range B15:P15 (the data row ITSELF has a merge across all columns except A)
- B16:P16 also merged
- Data in row 15: serial=1 in col A, rest in merged B15:P15
- This is the actual template the user's tool works with.
"""
import sys
import os
import shutil
import tempfile
import openpyxl
from copy import copy
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from services.chronology_generator.excel_writer import ChronologyExcelWriter
from services.chronology_generator.models import ChronologyEntry, ProjectChronology

def print_worksheet_state(ws, header_row, label):
    """Print complete worksheet state for debugging."""
    print(f"\n{'=' * 90}")
    print(f"  STATE: {label}")
    print(f"{'=' * 90}")
    print(f"  max_row={ws.max_row}, max_col={ws.max_column}, header_row={header_row}")
    
    print(f"\n  Merged ranges ({len(list(ws.merged_cells.ranges))}):")
    for mr in ws.merged_cells.ranges:
        master_val = ws.cell(row=mr.min_row, column=mr.min_col).value
        print(f"    {str(mr):20s}  master_val={repr(master_val)[:60]}")
    
    # Print data rows
    for r in range(header_row, min(ws.max_row + 1, header_row + 10)):
        print(f"\n  Row {r}:")
        for c in range(1, 17):  # Cols A through P
            val = ws.cell(row=r, column=c).value
            # Check if in merged range
            in_merge = ""
            for mr in ws.merged_cells.ranges:
                if mr.min_row <= r <= mr.max_row and mr.min_col <= c <= mr.max_col:
                    is_master = (r == mr.min_row and c == mr.min_col)
                    in_merge = " [MASTER]" if is_master else f" [MERGED->({mr.min_row},{mr.min_col})]"
                    break
            val_repr = repr(val)[:50] if val is not None else "<<EMPTY>>"
            print(f"    Col {c:2d} = {val_repr}{in_merge}")

def create_real_template(path):
    """Create a template that exactly matches the real Ladder_Chronology.xlsx structure."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Product Name ___Chronology"
    
    # Row 6: Merged A6:J6 like real template
    ws.merge_cells('A6:J6')
    ws['A6'] = 'Applicable Products:'
    
    # Row 14: Header row (exactly as in the real file)
    headers = {
        1: 'Serial No. ',
        2: 'Source Code Path',
        3: 'Bin File Name',
        4: 'CRC',
        5: 'Ladder Version No.',
        6: 'Reason For Upgrade',
        7: 'Testing Stage',
        8: 'PLC Model',
        9: 'Selpro Version & Path ',
        10: 'Bootloader Version',
        11: 'Release Date',
        12: 'Released By',
        13: 'Ladder  Release \u2013 To Production',
        14: 'Operator Procedure Modification',
        15: 'Automation Set Up Modification',
        16: 'Tested By',
    }
    for col, val in headers.items():
        ws.cell(row=14, column=col, value=val)
    
    # Row 15: This is the FIRST data row, but in the real template it has:
    #   - Col A (1): serial no = 1 (NOT merged)
    #   - Cols B-P (2-16): MERGED as B15:P15 with the source code path value
    # This is exactly the structure that causes the bug!
    ws.cell(row=15, column=1, value=1)  # Serial not merged
    ws.cell(row=15, column=2, value=r'C:\Users\SM 464\Desktop\Operational Package\2. Bin File\Master\Initial\FW_MIBRX-4M_BL20_V1.00.bin')
    ws.merge_cells('B15:P15')  # Merge AFTER writing value to master cell
    
    # Row 16: Merged B16:P16 (another merged row below)
    ws.cell(row=16, column=2, value=r'C:\Users\SM 464\Desktop\Operational Package\Source Code\Path')
    ws.merge_cells('B16:P16')
    
    # Add borders to data row 15
    from openpyxl.styles import Border, Side, Alignment
    thin = Side(style='thin')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    align = Alignment(horizontal='center')
    for c in range(1, 17):
        ws.cell(row=15, column=c).border = border
        ws.cell(row=15, column=c).alignment = align
    
    wb.save(path)
    wb.close()
    print(f"Created real-structure template at {path}")


def main():
    tmp_dir = tempfile.mkdtemp(prefix="debug_real_template_")
    print(f"Temp dir: {tmp_dir}")
    
    template_path = os.path.join(tmp_dir, "Ladder_Chronology.xlsx")
    output_path = os.path.join(tmp_dir, "Ladder_Chronology.xlsx")  # Same file = existing workbook
    
    # === STEP 1: Create the template exactly matching real structure ===
    create_real_template(template_path)
    
    # === STEP 2: Inspect what the template looks like ===
    wb = openpyxl.load_workbook(template_path)
    ws = wb.active
    print_worksheet_state(ws, 14, "TEMPLATE AS CREATED")
    wb.close()
    
    # === STEP 3: Run the ChronologyExcelWriter against this template ===
    print("\n\n" + "#" * 90)
    print("# RUNNING ChronologyExcelWriter.write_entries()")
    print("#" * 90)
    
    entry = ChronologyEntry(
        version="V1.01",
        source_code_path=Path(r"C:\Project\2. Bin File\Master\Initial\FW_MIBRX-4M_BL20_V1.01.bin"),
        bin_file_path=Path(r"C:\Project\2. Bin File\Master\Initial\FW_MIBRX-4M_BL20_V1.01.bin"),
        sdoc_file_path=None,
        bin_file_name="FW_MIBRX-4M_BL20_V1.01.bin",
        sdoc_file_name=None,
        crc="B4C82D9F",
        plc_model="MIBRX-4M",
        testing_stage="Master Initial",
        bootloader_version="BL20",
        selpro_version="SP4.0 C:/Selpro",
        reason_for_upgrade="Bug Fix Release",
        release_date="05/08/2026",
        released_by="Engineer B",
        ladder_release_to_production="Yes",
        operator_procedure_modification="No",
        automation_setup_modification="No",
        tested_by="QA Team B",
    )
    
    chronology = ProjectChronology(
        project_root=Path(r"C:\Project"),
        entries=[entry],
    )
    
    writer = ChronologyExcelWriter(template_path=Path(template_path))
    writer.write(chronology, Path(output_path))
    
    # === STEP 4: Inspect the result ===
    wb = openpyxl.load_workbook(output_path)
    ws = wb.active
    print_worksheet_state(ws, 14, "AFTER write_entries()")
    
    # === STEP 5: Check for corruption ===
    print("\n\n" + "!" * 90)
    print("  CORRUPTION CHECK")
    print("!" * 90)
    
    # Check row 15 (should be the NEW entry V1.01)
    print("\n  Row 15 (expected: V1.01 new entry):")
    for c in range(1, 17):
        val = ws.cell(row=15, column=c).value
        in_merge = ""
        for mr in ws.merged_cells.ranges:
            if mr.min_row <= 15 <= mr.max_row and mr.min_col <= c <= mr.max_col:
                in_merge = " [STILL MERGED!]"
                break
        val_repr = repr(val) if val is not None else "<<EMPTY>>"
        expected_empty = val is None or val == "" or val == 0
        status = "EMPTY!" if expected_empty and c > 1 else "OK"
        print(f"    Col {c:2d}: {val_repr:50s} {in_merge} {status}")
    
    # Check row 16 (should be the OLD entry V1.00 -- shifted down)
    print("\n  Row 16 (expected: V1.00 old entry):")
    for c in range(1, 17):
        val = ws.cell(row=16, column=c).value
        in_merge = ""
        for mr in ws.merged_cells.ranges:
            if mr.min_row <= 16 <= mr.max_row and mr.min_col <= c <= mr.max_col:
                in_merge = " [STILL MERGED!]"
                break
        val_repr = repr(val) if val is not None else "<<EMPTY>>"
        expected_empty = val is None or val == "" or val == 0
        status = "CORRUPTED!" if expected_empty and c > 1 else "OK"
        print(f"    Col {c:2d}: {val_repr:50s} {in_merge} {status}")
    
    # Row 17
    print("\n  Row 17 (may have old merged data):")
    for c in range(1, 17):
        val = ws.cell(row=17, column=c).value
        in_merge = ""
        for mr in ws.merged_cells.ranges:
            if mr.min_row <= 17 <= mr.max_row and mr.min_col <= c <= mr.max_col:
                in_merge = " [MERGED]"
                break
        val_repr = repr(val) if val is not None else "<<EMPTY>>"
        print(f"    Col {c:2d}: {val_repr:50s} {in_merge}")
    
    wb.close()
    
    # Cleanup
    shutil.rmtree(tmp_dir, ignore_errors=True)


if __name__ == "__main__":
    main()
