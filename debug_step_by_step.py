import sys
import os
import shutil
import tempfile
import openpyxl
from copy import copy
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from services.chronology_generator.excel_writer import ChronologyExcelWriter
from services.chronology_generator.models import ChronologyEntry, ProjectChronology

def print_state(ws, label):
    print(f"\n{'=' * 80}")
    print(f"--- {label} ---")
    print(f"{'=' * 80}")
    
    merged = list(ws.merged_cells.ranges)
    print(f"Merged ranges ({len(merged)}):")
    for mr in merged:
        print(f"  {str(mr)}")
    
    print("\nData Rows:")
    for r in range(14, 19):
        cols = []
        for c in range(1, 11): # Just first 10 columns for brevity
            val = ws.cell(row=r, column=c).value
            # Check if merged
            in_merge = ""
            for mr in ws.merged_cells.ranges:
                if mr.min_row <= r <= mr.max_row and mr.min_col <= c <= mr.max_col:
                    in_merge = " [M]"
                    break
            val_str = str(val)[:15] if val is not None else ""
            cols.append(f"{val_str:15s}{in_merge}")
        print(f"Row {r:2d}: | " + " | ".join(cols) + " |")

def main():
    tmp_dir = tempfile.mkdtemp(prefix="debug_step_by_step_")
    template_path = os.path.join(tmp_dir, "Template.xlsx")
    
    wb = openpyxl.Workbook()
    ws = wb.active
    
    # Setup Header (Row 14)
    headers = {
        1: 'Serial No.', 2: 'Source Code Path', 3: 'Bin File Name', 4: 'CRC', 
        5: 'Ladder Version No.', 6: 'Reason For Upgrade', 7: 'Testing Stage', 
        8: 'PLC Model', 9: 'Selpro Version & Path', 10: 'Bootloader Version',
        11: 'Release Date', 12: 'Released By', 13: 'Ladder Release-To Production',
        14: 'Operator Procedure Modification', 15: 'Automation Set Up Modification', 16: 'Tested By'
    }
    for col, val in headers.items():
        ws.cell(row=14, column=col, value=val)
        
    # Setup Row 15: Existing V1.00 data (UNMERGED)
    ws.cell(row=15, column=1, value=1)
    ws.cell(row=15, column=2, value="C:\\Path\\To\\V1.00")
    ws.cell(row=15, column=3, value="FW_V1.00.bin")
    ws.cell(row=15, column=4, value="ABCD1234")
    ws.cell(row=15, column=5, value="V1.00")
    ws.cell(row=15, column=6, value="Initial Release")
    ws.cell(row=15, column=7, value="Production")
    ws.cell(row=15, column=8, value="MIBRX-4M")
    
    # Setup Row 16: Instruction Row (MERGED B16:P16)
    ws.cell(row=16, column=2, value="Some instruction text here")
    ws.merge_cells('B16:P16')
    
    wb.save(template_path)
    wb.close()
    
    print_state(openpyxl.load_workbook(template_path).active, "1. Existing workbook loaded")
    
    # Now simulate the steps in write()
    writer = ChronologyExcelWriter(template_path=Path(template_path))
    wb = openpyxl.load_workbook(template_path)
    sheet = wb.active
    
    header_map = writer._build_header_map(sheet, 14)
    history_row = 15
    
    print("\n2. History row detected:", history_row)
    
    # 3. insert_rows
    sheet.insert_rows(history_row, amount=1)
    print_state(sheet, "3. Existing data after insert_rows() (and 4, 5)")
    
    # 6. _copy_row_format
    writer._copy_row_format(sheet, history_row + 1, history_row)
    print_state(sheet, "6. After _copy_row_format()")
    
    # 7. _write_entry
    entry = ChronologyEntry(
        version="V1.01",
        source_code_path=Path(r"C:\Path\To\V1.01"),
        bin_file_path=Path(r"C:\Path\To\V1.01\FW.bin"),
        sdoc_file_path=None,
        bin_file_name="FW_V1.01.bin",
        sdoc_file_name=None,
        crc="9876FEDC",
        plc_model="MIBRX-4M",
        testing_stage="Production",
        bootloader_version="BL20",
        selpro_version="SP4",
        reason_for_upgrade="Bug Fix",
        release_date="05/08/2026",
    )
    writer._write_entry(sheet, history_row, entry, header_map, 1, "05/08/2026")
    print_state(sheet, "7. After _write_entry()")
    
    # 8. _renumber_serials
    writer._renumber_serials(sheet, history_row, header_map)
    print_state(sheet, "8. After _renumber_serials()")
    
    # 9. Save
    output_path = os.path.join(tmp_dir, "Output.xlsx")
    writer._save(wb, Path(output_path))
    
    wb2 = openpyxl.load_workbook(output_path)
    print_state(wb2.active, "9. After Workbook Save & Reload")
    wb2.close()

if __name__ == "__main__":
    main()
