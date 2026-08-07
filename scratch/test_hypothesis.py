import openpyxl
from pathlib import Path

def test_template():
    # Find the original template
    template_paths = list(Path("C:/Users/SM 464/Desktop/Operational Package Validator/ProjectValidator/tests").rglob("base_template.xlsx"))
    if not template_paths:
        print("Template not found!")
        return
        
    wb = openpyxl.load_workbook(template_paths[0])
    ws = wb.active
    
    print(f"Original A5: {ws['A5'].value}")
    print(f"Original B5: {ws['B5'].value}")
    
    # Try merging B5:H5
    ws.merge_cells(start_row=5, start_column=2, end_row=5, end_column=8)
    
    ws['B5'].value = 'LPRS480'
    
    print(f"After merge A5: {ws['A5'].value}")
    print(f"After merge B5: {ws['B5'].value}")
    
    wb.save("scratch/test_overlap.xlsx")
    
    wb2 = openpyxl.load_workbook("scratch/test_overlap.xlsx")
    ws2 = wb2.active
    print(f"After save A5: {ws2['A5'].value}")
    print(f"After save B5: {ws2['B5'].value}")

test_template()
