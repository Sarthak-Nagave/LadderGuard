import openpyxl
from copy import copy
from dataclasses import dataclass

@dataclass
class Info:
    mode: str
    series_name: str
    products: list
    single_product: str
    parse_status: str = "Success"

class Writer:
    def _get_writable_cell(self, sheet, row, col):
        from openpyxl.worksheet.merge import MergedCell
        cell = sheet.cell(row=row, column=col)
        if isinstance(cell, MergedCell):
            for m_range in sheet.merged_cells.ranges:
                if row >= m_range.min_row and row <= m_range.max_row and \
                   col >= m_range.min_col and col <= m_range.max_col:
                    return sheet.cell(row=m_range.min_row, column=m_range.min_col)
        return cell
        
    def _set_cell_bold(self, sheet, row, col):
        pass
        
    def safe_val(self, val):
        return str(val) if val else ""
        
    def _strip_field_prefix(self, val, prefix):
        return val
        
    def _normalize_applicable_products(self, products):
        return products
        
    def _copy_row_format(self, sheet, src, dst):
        pass
        
    def _shift_row_breaks(self, sheet, src, dst):
        pass

    def run(self):
        wb = openpyxl.Workbook()
        sheet = wb.active
        
        info = Info(mode="Series", series_name="LPRS480", products=["Prod1", "Prod2"], single_product="")
        header_row_idx = 10
        
        product_label_row = 5
        product_label_col = 1
        product_value_col = 2
        applicable_label_row = 6
        products_start_row = 7
        
        product_label_text = "Product Series:"
        
        # Merge B5:H5
        sheet.merge_cells(start_row=product_label_row, start_column=2, end_row=product_label_row, end_column=8)
        # Merge A6:H6
        sheet.merge_cells(start_row=applicable_label_row, start_column=1, end_row=applicable_label_row, end_column=8)
        
        def _format_label(row: int, col: int, text: str) -> None:
            cell = self._get_writable_cell(sheet, row, col)
            cell.value = text
            if cell.alignment:
                align = copy(cell.alignment)
            else:
                from openpyxl.styles import Alignment
                align = Alignment()
            align.wrap_text = False
            cell.alignment = align

        _format_label(product_label_row, product_label_col, product_label_text)
        product_target = self._get_writable_cell(sheet, product_label_row, product_value_col)
        product_row, product_col = product_target.row, product_target.column
        
        val_to_write = info.series_name
        
        product_target.value = val_to_write
        
        original_header_row = header_row_idx
        last_product_row = 6 + len(info.products)
        required_header_row = last_product_row + 3
        rows_to_insert = max(0, required_header_row - original_header_row)
        
        rows_inserted = 0
        if rows_to_insert > 0:
            insert_idx = original_header_row
            
            ranges_to_shift = []
            for merged_range in list(sheet.merged_cells.ranges):
                if merged_range.min_row >= insert_idx:
                    sheet.unmerge_cells(str(merged_range))
                    shifted = copy(merged_range)
                    shifted.shift(row_shift=rows_to_insert, col_shift=0)
                    ranges_to_shift.append(shifted)

            sheet.insert_rows(insert_idx, amount=rows_to_insert)
            
            for shifted_range in ranges_to_shift:
                sheet.merge_cells(str(shifted_range))
                
            rows_inserted = rows_to_insert
            
        new_header_row = original_header_row + rows_inserted
        
        for clear_row in range(product_label_row, new_header_row):
            for clear_col in range(1, sheet.max_column + 1):
                self._get_writable_cell(sheet, clear_row, clear_col).value = ""
                
        _format_label(product_label_row, product_label_col, product_label_text)
        _format_label(product_row, product_col, val_to_write)
        _format_label(applicable_label_row, 1, "Applicable Products:")
        self._get_writable_cell(sheet, applicable_label_row, 2).value = ""
        self._get_writable_cell(sheet, applicable_label_row, 3).value = ""
        
        wb.save("scratch/test_writer.xlsx")
        
        wb2 = openpyxl.load_workbook("scratch/test_writer.xlsx")
        ws2 = wb2.active
        print(f"B5 value: {ws2['B5'].value}")
        print(f"A6 value: {ws2['A6'].value}")

Writer().run()
