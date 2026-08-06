import sys
import openpyxl
from pathlib import Path
from datetime import datetime
from services.chronology_generator.excel_writer import ChronologyExcelWriter
from services.chronology_generator.models import ProjectChronology, ChronologyEntry
from services.chronology_generator.pdf_parser import TestReportParser
import fitz

def trace():
    print('=== 1. & 2. & 3. Scanner / ChronologyEntry ===')
    pdf_path = Path(r'scratch/dummy_project/4. Test Report/Master/Initial/Master-Initial-Test-Report-sgn.pdf')
    entry = ChronologyEntry(
        source_code_path=Path('dummy'),
        bin_file_path=Path('dummy.bin'),
        sdoc_file_path=None,
        bin_file_name='dummy.bin',
        sdoc_file_name=None,
        version='V1.0',
        plc_model='Master',
        selpro_version='V1',
        bootloader_version='V1',
        crc='1234',
        testing_stage='Initial',
        release_date='01/01/2026',
        reason_for_upgrade='New feature',
        test_report_path=pdf_path
    )
    print('Firmware:', entry.testing_stage)
    print('Selected Test Report:', entry.test_report_path.name)
    print('Absolute path:', entry.test_report_path.absolute())
    print('entry.test_report_path:', entry.test_report_path)

    print('\n=== 4. Verify pdf_parser.py ===')
    parser = TestReportParser()
    info = parser.extract_product_info(entry.test_report_path)
    
    doc = fitz.open(entry.test_report_path)
    text = ''.join(p.get_text() for p in doc)
    print('PDF path:', entry.test_report_path)
    print('Pages scanned:', len(doc))
    print('Raw extracted text (first 1000 chars):', repr(text[:1000]))
    print('Detected mode:', info.mode)
    print('Series:', info.series_name)
    print('Single:', info.single_product)
    print('Products:', info.products)
    print('Parse status:', info.parse_status)

    print('\n=== 5. & 6. & 7. Verify excel_writer.py ===')
    template_path = Path(r'C:\Users\SM 464\Desktop\Operational Package Validator\ProjectValidator\backup_release_cleanup\archive\test_firmware_project\Master\Initial\Ladder_Chronology.xlsx')
    
    out_file = Path('scratch/debug_out.xlsx')
    import shutil
    shutil.copy(template_path, out_file)
    
    writer = ChronologyExcelWriter(template_path=template_path)
    chrono = ProjectChronology(project_root=Path('dummy'))
    chrono.add_entry(entry)
    
    # We will hook into writer logger to see what it prints
    import logging
    logging.basicConfig(level=logging.INFO, stream=sys.stdout)
    
    writer.write(chrono, out_file)

    wb = openpyxl.load_workbook(out_file)
    ws = wb.active
    product_row = None
    product_col = None
    applicable_row = None
    applicable_col = None

    for r in range(1, 40):
        for c in range(1, 10):
            val = writer.safe_val(ws.cell(row=r, column=c).value).lower()
            if not val:
                continue
            if val in ["product", "product:", "product :", "product series", "product series:", "product series :"]:
                product_row, product_col = r, c
            elif "product series" in val or "product :" in val or "product:" in val or val == "product":
                product_row, product_col = r, c
            if "applicable products" in val:
                applicable_row, applicable_col = r, c
    
    print(f'\nVerify worksheet search:')
    print(f'Product cell found: {product_row}, {product_col}')
    print(f'Applicable Products cell found: {applicable_row}, {applicable_col}')
    
    if product_row:
        val_written = writer.safe_val(ws.cell(row=product_row, column=product_col + 1).value)
        print(f'Value in cell next to product ({product_row}, {product_col+1}): {repr(val_written)}')
    else:
        print('NO PRODUCT ROW FOUND!')

if __name__ == '__main__':
    trace()
