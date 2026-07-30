"""
services.chronology_generator.template_loader
=============================================

Safely loads and validates Ladder Chronology Excel templates.

This module is strictly responsible for opening existing templates via openpyxl
and verifying the presence of all required column headers without modifying the workbook.

Author:
    Selec Controls Pvt. Ltd. - R&D

Python:
    3.12+
"""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from services.chronology_generator.exceptions import ChronologyTemplateError


class TemplateLoader:
    """
    Safely loads and validates Excel templates for the chronology generator.
    """

    # Definitive set of required normalized column headers
    REQUIRED_HEADERS = {
        "serial no.",
        "source code path",
        "bin file name",
        "crc",
        "ladder version no.",
        "reason for upgrade",
        "testing stage",
        "plc model",
        "selpro version & path",
        "bootloader version",
        "release date",
        "released by",
        "ladder release-to production",
        "operator procedure modification",
        "automation set up modification",
        "tested by",
    }

    def __init__(self) -> None:
        """
        Initialize the TemplateLoader.
        """
        self._logger = logging.getLogger(__name__)

    def load(self, template_path: Path) -> Workbook:
        """
        Safely loads an Excel template from the given path and validates its headers.

        Args:
            template_path: The file path to the template workbook.

        Returns:
            The loaded openpyxl Workbook instance.

        Raises:
            FileNotFoundError: If the template file does not exist.
            ChronologyTemplateError: If the file cannot be loaded or validation fails.
        """
        self._logger.info(f"Loading template from: {template_path}")

        if not template_path.exists():
            error_msg = f"Template file not found at path: {template_path}"
            self._logger.error(error_msg)
            raise FileNotFoundError(error_msg)

        try:
            # Load workbook read_only=False to support structure checks if needed,
            # but keep it light. We do not modify the workbook.
            workbook = openpyxl.load_workbook(filename=template_path, data_only=True)
        except Exception as exc:
            error_msg = f"Failed to open workbook at '{template_path}': {exc}"
            self._logger.error(error_msg)
            raise ChronologyTemplateError(message=error_msg, original_exception=exc) from exc

        # Perform header validation immediately upon load
        self.validate(workbook)

        self._logger.info("Template successfully loaded and validated.")
        return workbook

    def validate(self, workbook: Workbook) -> None:
        """
        Validates the active worksheet of the workbook for required headers.

        Locates the header row dynamically and checks against the required header list.

        Args:
            workbook: The loaded openpyxl Workbook to validate.

        Raises:
            ChronologyTemplateError: If no active sheet is found, the header row is missing,
                                      or any required headers are absent.
        """
        self._logger.debug("Validating workbook headers.")

        sheet = workbook.active
        if sheet is None:
            raise ChronologyTemplateError("The template workbook contains no active worksheet.")

        header_row_idx = self._find_header_row(sheet)
        self._verify_required_headers(sheet, header_row_idx)

    def _find_header_row(self, sheet: Worksheet) -> int:
        """
        Dynamically scans the sheet to locate the row containing the chronology headers.

        Searches for the 'Serial No.' identifier within the first 50 rows.

        Args:
            sheet: The worksheet being scanned.

        Returns:
            The 1-based row index of the header row.

        Raises:
            ChronologyTemplateError: If the header row cannot be located.
        """
        for row_idx in range(1, 51):
            for col_idx in range(1, sheet.max_column + 1):
                cell_value = sheet.cell(row=row_idx, column=col_idx).value
                if self._normalize_header(cell_value) == "serial no.":
                    self._logger.debug(f"Located header row at index: {row_idx}")
                    return row_idx

        error_msg = "Could not locate the header row (missing 'Serial No.' column) in the template."
        self._logger.error(error_msg)
        raise ChronologyTemplateError(error_msg)

    def _verify_required_headers(self, sheet: Worksheet, header_row_idx: int) -> None:
        """
        Verifies that all required headers are present in the identified header row.

        Args:
            sheet: The worksheet containing the headers.
            header_row_idx: The 1-based index of the header row.

        Raises:
            ChronologyTemplateError: If any required headers are missing.
        """
        found_headers: set[str] = set()

        for col_idx in range(1, sheet.max_column + 1):
            cell_value = sheet.cell(row=header_row_idx, column=col_idx).value
            normalized = self._normalize_header(cell_value)
            if normalized:
                found_headers.add(normalized)

        missing_headers = self.REQUIRED_HEADERS - found_headers
        if missing_headers:
            missing_str = ", ".join(f"'{h}'" for h in sorted(missing_headers))
            error_msg = f"Template is missing required column headers: {missing_str}"
            self._logger.error(error_msg)
            raise ChronologyTemplateError(error_msg)

        self._logger.debug("All required headers successfully verified.")

    @staticmethod
    def _normalize_header(value: Any) -> str:
        """
        Normalizes a header cell value for robust matching.

        - Replaces newlines with spaces.
        - Replaces Unicode dashes with standard '-'.
        - Converts to lowercase.
        - Removes spaces around hyphens.
        - Collapses multiple spaces into one.
        - Strips leading/trailing spaces.

        Args:
            value: The raw cell value.

        Returns:
            The normalized header string.
        """
        if value is None:
            return ""
        
        text = str(value)
        # Replace newlines/carriage returns with spaces
        text = re.sub(r"[\r\n]+", " ", text)
        # Replace unicode dashes with standard hyphen
        text = re.sub(r"[\u2010\u2011\u2012\u2013\u2014\u2015\u2212]", "-", text)
        # Convert to lowercase
        text = text.casefold()
        # Remove spaces around hyphens (e.g., " - " -> "-")
        text = re.sub(r"\s*-\s*", "-", text)
        # Collapse multiple spaces into one and strip
        return re.sub(r"\s+", " ", text).strip()