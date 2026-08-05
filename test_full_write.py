import openpyxl
from pathlib import Path
from services.chronology_generator.models import ProjectChronology, ChronologyEntry
from services.chronology_generator.excel_writer import ChronologyExcelWriter

def main():
    template_path = Path("template_test.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    
    headers = {
        1: 'Serial No.', 2: 'Source Code Path', 3: 'Bin File Name', 4: 'CRC', 
        5: 'Ladder Version No.', 6: 'Reason For Upgrade', 7: 'Testing Stage', 
        8: 'PLC Model', 9: 'Selpro Version & Path', 10: 'Bootloader Version',
        11: 'Release Date', 12: 'Released By', 13: 'Ladder Release-To Production',
        14: 'Operator Procedure Modification', 15: 'Automation Set Up Modification', 16: 'Tested By'
    }
    for col, val in headers.items():
        ws.cell(row=14, column=col, value=val)
        
    data = {
        1: 1, 2: r"C:\Path\To\V1.0", 3: "FW_V1.00.bin", 4: "ABCD1234", 
        5: "V1.00", 6: "Initial Release", 7: "Production", 8: "MIBRX-4M"
    }
    for col, val in data.items():
        ws.cell(row=15, column=col, value=val)
        
    ws.cell(row=16, column=2, value="Some instructions")
    ws.merge_cells("B16:P16")
    
    # Add a footer to test preservation
    ws.cell(row=20, column=2, value="Footer Text")
    ws.merge_cells("B20:D20")
    
    wb.save(template_path)
    wb.close()
    
    chronology = ProjectChronology(Path("."))
    entry = ChronologyEntry(
        version="V1.01",
        source_code_path=Path(r"C:\Path\To\V1.01"),
        bin_file_path=Path(r"C:\Path\To\V1.01\FW.bin"),
        sdoc_file_path=None,
        bin_file_name="FW_V1.01.bin",
        sdoc_file_name=None,
        crc="9876FEDC",
        plc_model="MIBRX-4M",
        testing_stage="Production",
        bootloader_version="BL20",
        selpro_version="SP4",
        reason_for_upgrade="Bug Fix",
        release_date="05/08/2026",
    )
    chronology.entries.append(entry)
    
    writer = ChronologyExcelWriter(template_path)
    writer.write(chronology, template_path)
    
    # Reload and print
    wb = openpyxl.load_workbook(template_path)
    ws = wb.active
    print("Merged cells:", list(ws.merged_cells.ranges))
    for r in range(14, 23):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, 6)]
        print(f"Row {r}: {row_vals}")

if __name__ == "__main__":
    main()
