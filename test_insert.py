import openpyxl
from openpyxl.workbook import Workbook

wb = Workbook()
ws = wb.active
ws.cell(row=10, column=1, value='header1')
ws.cell(row=10, column=2, value='header2')

ws.cell(row=11, column=1, value='val1')
ws.merge_cells('A11:B11')

ws.cell(row=12, column=1, value='val2')
ws.merge_cells('A12:B12')

ws.cell(row=15, column=1, value='footer')
ws.merge_cells('A15:B15')

print('Before insert:')
for r in ws.merged_cells.ranges: print(r)

ws.insert_rows(11, 1)

print('After insert:')
for r in ws.merged_cells.ranges: print(r)
