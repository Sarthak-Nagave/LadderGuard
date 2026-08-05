import openpyxl
wb = openpyxl.Workbook()
ws = wb.active
ws.cell(row=15, column=2, value="Data")
ws.merge_cells('B15:P15')
ws.insert_rows(15, 1)
print(list(ws.merged_cells.ranges))
