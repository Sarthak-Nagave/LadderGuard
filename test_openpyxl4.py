import openpyxl
from copy import copy
wb = openpyxl.Workbook()
ws = wb.active
ws.cell(row=15, column=2, value="Data")
ws.cell(row=16, column=2, value="Instruction")
ws.merge_cells('B16:P16')

# Try manually shifting
ranges_to_shift = []
for mr in list(ws.merged_cells.ranges):
    if mr.min_row >= 15:
        ws.merged_cells.remove(mr)
        mr.shift(row_shift=1, col_shift=0)
        ranges_to_shift.append(mr)

ws.insert_rows(15, 1)

for mr in ranges_to_shift:
    ws.merged_cells.append(mr)
    
print(list(ws.merged_cells.ranges))
