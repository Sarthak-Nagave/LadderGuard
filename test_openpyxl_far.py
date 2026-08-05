import openpyxl
wb = openpyxl.Workbook()
ws = wb.active
ws.cell(row=5, column=2, value="Footer")
ws.merge_cells('B5:D5')
ws.insert_rows(2, 1)
print(list(ws.merged_cells.ranges))
