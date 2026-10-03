from __future__ import annotations

import os
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import openpyxl
from PySide6.QtWidgets import QApplication

from gui.chronology.chronology_dialog import ChronologyDialog
from services.chronology_generator.generator import ChronologyGenerator
from services.chronology_generator.models import ChronologyEntry, ProjectChronology
from services.chronology_generator.project_scanner import ProjectScanner


class ChronologyGenerationWorkflowTests(unittest.TestCase):
    def _create_template(self, path: Path) -> None:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        headers = [
            "Serial No.",
            "Source Code Path",
            "Bin File Name",
            "CRC",
            "Ladder Version No.",
            "Reason for Upgrade",
            "Testing Stage",
            "PLC Model",
            "Selpro Version & Path",
            "Bootloader Version",
            "Release Date",
            "Released By",
            "Ladder Release-To Production",
            "Operator Procedure Modification",
            "Automation Set Up Modification",
            "Tested By",
        ]
        for idx, header in enumerate(headers, start=1):
            sheet.cell(row=1, column=idx).value = header
        workbook.save(path)

    def _header_map(self, sheet) -> dict[str, int]:
        mapping: dict[str, int] = {}
        for col_idx in range(1, sheet.max_column + 1):
            value = sheet.cell(row=1, column=col_idx).value
            if value is None:
                continue
            mapping[str(value).strip().casefold()] = col_idx
        return mapping

    def test_scanner_selects_newest_bin_and_creates_one_entry_per_stage(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            stage = root / "2. Bin File" / "Master" / "Initial"
            stage.mkdir(parents=True, exist_ok=True)

            old_bin = stage / "FW_MIBRX-4M_BL20_V1.00.bin"
            new_bin = stage / "FW_MIBRX-4M_BL20_V1.01.bin"
            old_bin.write_bytes(b"old")
            new_bin.write_bytes(b"new")

            # Ensure deterministic newest selection by modified time.
            os.utime(old_bin, (1_700_000_000, 1_700_000_000))
            os.utime(new_bin, (1_800_000_000, 1_800_000_000))

            scanner = ProjectScanner(root)
            chronology = scanner.scan()

            self.assertEqual(len(chronology.entries), 1)
            entry = chronology.entries[0]
            self.assertEqual(entry.bin_file_name, "FW_MIBRX-4M_BL20_V1.01.bin")
            self.assertEqual(entry.version, "V1.01")
            self.assertEqual(entry.plc_model, "MIBRX-4M")
            self.assertEqual(entry.bootloader_version, "BL20")
            self.assertEqual(entry.testing_stage, "Master Initial")
            self.assertEqual(entry.source_code_path, new_bin.resolve())

    def test_scanner_searches_bins_only_in_discovered_folder_not_subfolders(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            stage = root / "2. Bin File" / "Master" / "Initial"
            backup = stage / "Backup"
            stage.mkdir(parents=True, exist_ok=True)
            backup.mkdir(parents=True, exist_ok=True)

            selected_bin = stage / "FW_MIBRX-4M_BL20_V1.01.bin"
            ignored_backup_bin = backup / "FW_MIBRX-4M_BL20_V9.99.bin"
            selected_bin.write_bytes(b"main")
            ignored_backup_bin.write_bytes(b"backup")

            os.utime(selected_bin, (1_700_000_000, 1_700_000_000))
            os.utime(ignored_backup_bin, (1_900_000_000, 1_900_000_000))

            scanner = ProjectScanner(root)
            chronology = scanner.scan()

            self.assertEqual(len(chronology.entries), 1)
            self.assertEqual(chronology.entries[0].bin_file_name, "FW_MIBRX-4M_BL20_V1.01.bin")

    def test_scanner_uses_filename_tiebreak_when_modified_time_identical(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            stage = root / "2. Bin File" / "Master" / "Initial"
            stage.mkdir(parents=True, exist_ok=True)

            bin_a = stage / "FW_MIBRX-4M_BL20_V1.01_A.bin"
            bin_b = stage / "FW_MIBRX-4M_BL20_V1.01_B.bin"
            bin_a.write_bytes(b"a")
            bin_b.write_bytes(b"b")

            same_mtime = 1_800_000_000
            os.utime(bin_a, (same_mtime, same_mtime))
            os.utime(bin_b, (same_mtime, same_mtime))

            scanner = ProjectScanner(root)
            chronology = scanner.scan()

            self.assertEqual(len(chronology.entries), 1)
            self.assertEqual(chronology.entries[0].bin_file_name, "FW_MIBRX-4M_BL20_V1.01_A.bin")

    def test_scanner_metadata_comes_from_selected_newest_bin_only(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            stage = root / "2. Bin File" / "Master" / "Initial"
            stage.mkdir(parents=True, exist_ok=True)

            older_bin = stage / "FW_SLV_FINAL_MIBRX-4M_BL10_V1.00.bin"
            newer_bin = stage / "FW_MST_INITIAL_FLEXYS-2M_BL30_V2.50.bin"
            older_bin.write_bytes(b"old")
            newer_bin.write_bytes(b"new")

            os.utime(older_bin, (1_700_000_000, 1_700_000_000))
            os.utime(newer_bin, (1_800_000_000, 1_800_000_000))

            scanner = ProjectScanner(root)
            chronology = scanner.scan()

            self.assertEqual(len(chronology.entries), 1)
            entry = chronology.entries[0]

            self.assertEqual(entry.bin_file_name, "FW_MST_INITIAL_FLEXYS-2M_BL30_V2.50.bin")
            self.assertEqual(entry.version, "V2.50")
            self.assertEqual(entry.plc_model, "FLEXYS-2M")
            self.assertEqual(entry.bootloader_version, "BL30")
            self.assertEqual(entry.testing_stage, "Master Initial")
            self.assertEqual(entry.source_code_path, newer_bin.resolve())

    def test_generator_creates_new_stage_chronology_from_template(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            template_path = root / "template.xlsx"
            self._create_template(template_path)

            generator = ChronologyGenerator(root, template_path)
            selected_bin = root / "2. Bin File" / "Master" / "Initial" / "FW_MIBRX-4M_BL20_V1.00_MST_INITIAL.bin"
            selected_bin.parent.mkdir(parents=True, exist_ok=True)
            selected_bin.write_bytes(b"x")
            entry = ChronologyEntry(
                source_code_path=selected_bin.resolve(),
                bin_file_path=selected_bin,
                sdoc_file_path=None,
                bin_file_name="FW_MIBRX-4M_BL20_V1.00_MST_INITIAL.bin",
                sdoc_file_name=None,
                version="V1.00",
                plc_model="MIBRX-4M",
                selpro_version="",
                bootloader_version="BL20",
                crc="A3F91C7E",
                testing_stage="Master Initial",
                release_date="",
                reason_for_upgrade="",
            )
            chronology = ProjectChronology(project_root=root, entries=[entry])
            generator.scan_project = lambda: chronology

            generator.generate()

            output_path = root / "7. Chronology" / "Master" / "Initial" / "Ladder_Chronology.xlsx"
            self.assertTrue(output_path.exists())
            self.assertFalse((root / "2. Bin File" / "Master" / "Initial" / "Ladder_Chronology.xlsx").exists())

            workbook = openpyxl.load_workbook(output_path)
            sheet = workbook.active
            headers = self._header_map(sheet)

            self.assertEqual(sheet.cell(row=2, column=headers["serial no."]).value, 1)
            self.assertEqual(sheet.cell(row=2, column=headers["bin file name"]).value, "FW_MIBRX-4M_BL20_V1.00_MST_INITIAL.bin")
            self.assertEqual(sheet.cell(row=2, column=headers["ladder version no."]).value, "V1.00")
            self.assertEqual(sheet.cell(row=2, column=headers["plc model"]).value, "MIBRX-4M")
            self.assertEqual(sheet.cell(row=2, column=headers["bootloader version"]).value, "BL20")
            self.assertEqual(sheet.cell(row=2, column=headers["source code path"]).value, str(selected_bin.resolve()))
            self.assertEqual(sheet.cell(row=2, column=headers["crc"]).value, "A3F91C7E")
            self.assertEqual(
                sheet.cell(row=2, column=headers["release date"]).value,
                datetime.now().strftime("%d/%m/%Y"),
            )

    def test_generator_uses_existing_stage_workbook_and_inserts_newest_history_row(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            template_path = root / "template.xlsx"
            self._create_template(template_path)

            bin_folder = root / "2. Bin File" / "Master" / "Initial"
            bin_folder.mkdir(parents=True, exist_ok=True)
            output_path = root / "7. Chronology" / "Master" / "Initial" / "Ladder_Chronology.xlsx"
            output_path.parent.mkdir(parents=True, exist_ok=True)
            self._create_template(output_path)

            existing_wb = openpyxl.load_workbook(output_path)
            existing_sheet = existing_wb.active
            headers = self._header_map(existing_sheet)

            old_row = 2
            existing_sheet.cell(row=old_row, column=headers["serial no."]).value = 1
            existing_sheet.cell(row=old_row, column=headers["bin file name"]).value = "FW_MIBRX-4M_BL20_V1.00_MST_INITIAL.bin"
            existing_sheet.cell(row=old_row, column=headers["ladder version no."]).value = "V1.00"
            existing_sheet.cell(row=old_row, column=headers["testing stage"]).value = "Master Initial"
            existing_sheet.cell(row=old_row, column=headers["source code path"]).value = str(
                (bin_folder / "FW_MIBRX-4M_BL20_V1.00_MST_INITIAL.bin").resolve()
            )
            existing_sheet.cell(row=old_row, column=headers["plc model"]).value = "MIBRX-4M"
            existing_sheet.cell(row=old_row, column=headers["bootloader version"]).value = "BL20"
            existing_sheet.cell(row=old_row, column=headers["release date"]).value = "01/07/2026"
            existing_sheet.cell(row=old_row, column=headers["bin file name"]).font = openpyxl.styles.Font(bold=True)
            existing_wb.save(output_path)

            generator = ChronologyGenerator(root, template_path)
            selected_bin = bin_folder / "FW_MIBRX-4M_BL20_V1.01_MST_INITIAL.bin"
            selected_bin.write_bytes(b"new")
            new_entry = ChronologyEntry(
                source_code_path=selected_bin.resolve(),
                bin_file_path=selected_bin,
                sdoc_file_path=None,
                bin_file_name="FW_MIBRX-4M_BL20_V1.01_MST_INITIAL.bin",
                sdoc_file_name=None,
                version="V1.01",
                plc_model="MIBRX-4M",
                selpro_version="",
                bootloader_version="BL20",
                crc="",
                testing_stage="Master Initial",
                release_date="",
                reason_for_upgrade="",
            )
            chronology = ProjectChronology(project_root=root, entries=[new_entry])
            generator.scan_project = lambda: chronology

            generator.generate()

            workbook = openpyxl.load_workbook(output_path)
            sheet = workbook.active
            headers = self._header_map(sheet)

            self.assertEqual(sheet.cell(row=2, column=headers["serial no."]).value, 2)
            self.assertEqual(sheet.cell(row=2, column=headers["ladder version no."]).value, "V1.01")
            self.assertEqual(sheet.cell(row=3, column=headers["serial no."]).value, 1)
            self.assertEqual(sheet.cell(row=3, column=headers["ladder version no."]).value, "V1.00")
            self.assertEqual(sheet.cell(row=2, column=headers["source code path"]).value, str(selected_bin.resolve()))

            # Existing row formatting is preserved after insertion/shift.
            self.assertTrue(sheet.cell(row=3, column=headers["bin file name"]).font.bold)

    def test_generator_does_not_duplicate_top_release_when_same_bin_already_exists(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            template_path = root / "template.xlsx"
            self._create_template(template_path)

            bin_folder = root / "2. Bin File" / "Master" / "Initial"
            bin_folder.mkdir(parents=True, exist_ok=True)
            output_path = root / "7. Chronology" / "Master" / "Initial" / "Ladder_Chronology.xlsx"
            output_path.parent.mkdir(parents=True, exist_ok=True)
            self._create_template(output_path)

            selected_bin = bin_folder / "FW_MIBRX-4M_BL20_V1.01_MST_INITIAL.bin"
            selected_bin.write_bytes(b"new")

            existing_wb = openpyxl.load_workbook(output_path)
            existing_sheet = existing_wb.active
            headers = self._header_map(existing_sheet)

            existing_sheet.cell(row=2, column=headers["serial no."]).value = 1
            existing_sheet.cell(row=2, column=headers["bin file name"]).value = selected_bin.name
            existing_sheet.cell(row=2, column=headers["ladder version no."]).value = "V1.01"
            existing_sheet.cell(row=2, column=headers["testing stage"]).value = "Master Initial"
            existing_sheet.cell(row=2, column=headers["source code path"]).value = str(selected_bin.resolve())
            existing_sheet.cell(row=2, column=headers["plc model"]).value = "MIBRX-4M"
            existing_sheet.cell(row=2, column=headers["bootloader version"]).value = "BL20"
            existing_sheet.cell(row=2, column=headers["release date"]).value = "01/07/2026"

            existing_sheet.cell(row=3, column=headers["serial no."]).value = 2
            existing_sheet.cell(row=3, column=headers["bin file name"]).value = "FW_MIBRX-4M_BL20_V1.00_MST_INITIAL.bin"
            existing_sheet.cell(row=3, column=headers["ladder version no."]).value = "V1.00"
            existing_wb.save(output_path)

            generator = ChronologyGenerator(root, template_path)
            entry = ChronologyEntry(
                source_code_path=selected_bin.resolve(),
                bin_file_path=selected_bin,
                sdoc_file_path=None,
                bin_file_name=selected_bin.name,
                sdoc_file_name=None,
                version="V1.01",
                plc_model="MIBRX-4M",
                selpro_version="",
                bootloader_version="BL20",
                crc="",
                testing_stage="Master Initial",
                release_date="",
                reason_for_upgrade="",
            )
            chronology = ProjectChronology(project_root=root, entries=[entry])
            generator.scan_project = lambda: chronology

            generator.generate()

            workbook = openpyxl.load_workbook(output_path)
            sheet = workbook.active

            self.assertEqual(sheet.cell(row=2, column=headers["serial no."]).value, 2)
            self.assertEqual(sheet.cell(row=2, column=headers["bin file name"]).value, selected_bin.name)
            self.assertEqual(sheet.cell(row=3, column=headers["serial no."]).value, 1)
            self.assertEqual(sheet.max_row, 3)

    def test_dialog_builds_dynamic_cards_from_firmware_folders(self) -> None:
        QApplication.instance() or QApplication([])

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            stage1 = root / "2. Bin File" / "Master" / "Initial"
            stage2 = root / "2. Bin File" / "Slave" / "Final"
            stage3 = root / "2. Bin File" / "UUT" / "QC"
            stage1.mkdir(parents=True, exist_ok=True)
            stage2.mkdir(parents=True, exist_ok=True)
            stage3.mkdir(parents=True, exist_ok=True)

            bin1 = stage1 / "FW_MST_INITIAL_MIBRX-4M_BL20_V1.00.bin"
            bin2 = stage2 / "FW_SLV_FINAL_FLEXYS-2M_BL10_V1.01.bin"
            bin3 = stage3 / "FW_UUT_QC_MIBRX-2M_BL30_V1.02.bin"
            bin1.write_bytes(b"1")
            bin2.write_bytes(b"2")
            bin3.write_bytes(b"3")

            records = {
                "Master/Initial": {"status": "PASS", "bin_file": str(bin1), "bin_name": bin1.name, "bin_crc": "AAAABBBB"},
                "Slave/Final": {"status": "PASS", "bin_file": str(bin2), "bin_name": bin2.name, "bin_crc": "CCCCDDDD"},
                "UUT/QC": {"status": "PASS", "bin_file": str(bin3), "bin_name": bin3.name, "bin_crc": "EEEEFFFF"},
            }

            dialog = ChronologyDialog(project_folder=root, validation_bin_crc_records=records)
            self.assertEqual(len(dialog._entry_cards), 3)

            source_paths = {card["line_source_path"].text() for card in dialog._entry_cards}
            self.assertEqual(
                source_paths,
                {""},
            )

            folder_labels = {card["line_firmware_folder"].text() for card in dialog._entry_cards}
            self.assertEqual(folder_labels, {"Master/Initial", "Slave/Final", "UUT/QC"})

            for card in dialog._entry_cards:
                self.assertFalse(card["line_source_path"].isReadOnly())

            dialog.close()

    def test_dialog_generate_button_waits_for_active_card_required_fields(self) -> None:
        QApplication.instance() or QApplication([])

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            template_path = root / "template.xlsx"
            self._create_template(template_path)

            stage1 = root / "2. Bin File" / "Master" / "Initial"
            stage2 = root / "2. Bin File" / "Slave" / "Final"
            stage1.mkdir(parents=True, exist_ok=True)
            stage2.mkdir(parents=True, exist_ok=True)

            (stage1 / "FW_MST_INITIAL_MIBRX-4M_BL20_V1.00.bin").write_bytes(b"1")
            (stage2 / "FW_SLV_FINAL_FLEXYS-2M_BL10_V1.01.bin").write_bytes(b"2")

            records = {
                "Master/Initial": {"status": "PASS", "bin_file": str(stage1 / "FW_MST_INITIAL_MIBRX-4M_BL20_V1.00.bin"), "bin_name": "FW_MST_INITIAL_MIBRX-4M_BL20_V1.00.bin", "bin_crc": "AAAABBBB"},
                "Slave/Final": {"status": "PASS", "bin_file": str(stage2 / "FW_SLV_FINAL_FLEXYS-2M_BL10_V1.01.bin"), "bin_name": "FW_SLV_FINAL_FLEXYS-2M_BL10_V1.01.bin", "bin_crc": "CCCCDDDD"},
            }

            dialog = ChronologyDialog(project_folder=root, validation_bin_crc_records=records)
            dialog._update_generate_button_state()
            self.assertFalse(dialog.ui.btn_generate.isEnabled())

            # Find the active card
            active_widget = dialog._stacked_cards.currentWidget()
            active_card = next(c for c in dialog._entry_cards if c["group_box"] == active_widget)

            active_card["line_reason"].setText("Release update")
            active_card["line_released_by"].setText("Dev A")
            active_card["line_tested_by"].setText("QA A")
            active_card["combo_ladder_release"].setCurrentText("Yes")
            active_card["combo_operator_modification"].setCurrentText("No")
            active_card["combo_automation_modification"].setCurrentText("No")
            active_card["combo_prepared_by"].setCurrentText("John")
            active_card["combo_checked_by"].setCurrentText("Jane")
            active_card["combo_approved_by"].setCurrentText("Boss")
            dialog._update_generate_button_state()

            # Generate button should be enabled since active card is filled out
            self.assertTrue(dialog.ui.btn_generate.isEnabled())
            dialog.close()

    def test_dialog_user_edits_apply_to_matching_card_only(self) -> None:
        QApplication.instance() or QApplication([])

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            template_path = root / "template.xlsx"
            self._create_template(template_path)

            stage1 = root / "2. Bin File" / "Master" / "Initial"
            stage2 = root / "2. Bin File" / "Slave" / "Final"
            stage1.mkdir(parents=True, exist_ok=True)
            stage2.mkdir(parents=True, exist_ok=True)

            (stage1 / "FW_MST_INITIAL_MIBRX-4M_BL20_V1.00.bin").write_bytes(b"1")
            (stage2 / "FW_SLV_FINAL_FLEXYS-2M_BL10_V1.01.bin").write_bytes(b"2")

            records = {
                "Master/Initial": {"status": "PASS", "bin_file": str(stage1 / "FW_MST_INITIAL_MIBRX-4M_BL20_V1.00.bin"), "bin_name": "FW_MST_INITIAL_MIBRX-4M_BL20_V1.00.bin", "bin_crc": "AAAABBBB"},
                "Slave/Final": {"status": "PASS", "bin_file": str(stage2 / "FW_SLV_FINAL_FLEXYS-2M_BL10_V1.01.bin"), "bin_name": "FW_SLV_FINAL_FLEXYS-2M_BL10_V1.01.bin", "bin_crc": "CCCCDDDD"},
            }

            dialog = ChronologyDialog(project_folder=root, validation_bin_crc_records=records)

            card1 = dialog._entry_cards[0]
            card2 = dialog._entry_cards[1]

            card1["line_reason"].setText("Card1 Reason")
            card1["line_released_by"].setText("Card1 Dev")
            card1["line_tested_by"].setText("Card1 QA")
            card1["line_selpro_version"].setText("SP1")
            card1["line_selpro_path"].setText("C:/selpro/a")
            card1["combo_ladder_release"].setCurrentText("Yes")
            card1["combo_operator_modification"].setCurrentText("No")
            card1["combo_automation_modification"].setCurrentText("No")

            card2["line_reason"].setText("Card2 Reason")
            card2["line_released_by"].setText("Card2 Dev")
            card2["line_tested_by"].setText("Card2 QA")
            card2["line_selpro_version"].setText("SP2")
            card2["line_selpro_path"].setText("C:/selpro/b")
            card2["combo_ladder_release"].setCurrentText("No")
            card2["combo_operator_modification"].setCurrentText("Yes")
            card2["combo_automation_modification"].setCurrentText("Yes")

            dialog._update_chronology_models()

            entry1 = card1["entry"]
            entry2 = card2["entry"]

            self.assertEqual(entry1.reason_for_upgrade, "Card1 Reason")
            self.assertEqual(entry1.released_by, "Card1 Dev")
            self.assertEqual(entry1.tested_by, "Card1 QA")
            self.assertEqual(entry1.selpro_version, "SP1 C:/selpro/a")
            self.assertEqual(entry1.ladder_release_to_production, "Yes")
            self.assertEqual(entry1.operator_procedure_modification, "No")
            self.assertEqual(entry1.automation_setup_modification, "No")

            self.assertEqual(entry2.reason_for_upgrade, "Card2 Reason")
            self.assertEqual(entry2.released_by, "Card2 Dev")
            self.assertEqual(entry2.tested_by, "Card2 QA")
            self.assertEqual(entry2.selpro_version, "SP2 C:/selpro/b")
            self.assertEqual(entry2.ladder_release_to_production, "No")
            self.assertEqual(entry2.operator_procedure_modification, "Yes")
            self.assertEqual(entry2.automation_setup_modification, "Yes")
            dialog.close()

    def test_dialog_populates_crc_only_for_passed_validation_records(self) -> None:
        QApplication.instance() or QApplication([])

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            stage_pass = root / "2. Bin File" / "Master" / "Initial"
            stage_fail = root / "2. Bin File" / "Slave" / "Final"
            stage_pass.mkdir(parents=True, exist_ok=True)
            stage_fail.mkdir(parents=True, exist_ok=True)

            pass_bin = stage_pass / "FW_MST_INITIAL_MIBRX-4M_BL20_V1.00.bin"
            fail_bin = stage_fail / "FW_SLV_FINAL_FLEXYS-2M_BL10_V1.01.bin"
            pass_bin.write_bytes(b"pass")
            fail_bin.write_bytes(b"fail")

            records = {
                "Master/Initial": {
                    "status": "PASS",
                    "bin_file": str(pass_bin),
                    "bin_name": pass_bin.name,
                    "bin_crc": "A3F91C7E",
                },
                "Slave/Final": {
                    "status": "FAILED",
                    "bin_file": str(fail_bin),
                    "bin_name": fail_bin.name,
                    "bin_crc": "DEADBEEF",
                    "reason": "CRC Mismatch",
                },
            }

            dialog = ChronologyDialog(
                project_folder=root,
                validation_bin_crc_records=records,
            )

            card_map = {card["line_firmware_folder"].text(): card for card in dialog._entry_cards}
            self.assertEqual(len(card_map), 1)
            self.assertEqual(card_map["Master/Initial"]["entry"].crc, "A3F91C7E")
            self.assertEqual(dialog._skipped_firmware[0]["firmware_folder"], "Slave/Final")
            self.assertEqual(dialog._skipped_firmware[0]["reason"], "CRC Mismatch")
            dialog.close()


if __name__ == "__main__":
    unittest.main()

