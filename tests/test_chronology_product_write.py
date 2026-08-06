from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl

from services.chronology_generator.excel_writer import ChronologyExcelWriter
from services.chronology_generator.models import ChronologyEntry, ProductInfo, ProjectChronology


class ChronologyProductWriteTests(unittest.TestCase):
    def _create_template(self, path: Path) -> None:
        workbook = openpyxl.Workbook()
        sheet = workbook.active

        sheet["A3"] = "Product:"
        sheet.merge_cells("B3:D3")
        sheet["B3"] = "OLD_PRODUCT"

        sheet["A8"] = "Applicable Products"

        # Keep only three rows before history so Series mode with 4 products inserts one row.
        sheet["A9"] = 1
        sheet["B9"] = "LEGACY-1"
        sheet["A10"] = 2
        sheet["B10"] = "LEGACY-2"
        sheet["A11"] = 3
        sheet["B11"] = "LEGACY-3"

        header_row = 12
        headers = [
            "Serial No.",
            "Source Code Path",
            "Bin File Name",
            "CRC",
            "Ladder Version No.",
            "Reason for Upgrade",
            "Testing Stage",
            "PLC Model",
            "Selpro Version & Path",
            "Bootloader Version",
            "Release Date",
            "Released By",
            "Ladder Release-To Production",
            "Operator Procedure Modification",
            "Automation Set Up Modification",
            "Tested By",
        ]
        for idx, header in enumerate(headers, start=1):
            sheet.cell(row=header_row, column=idx).value = header

        workbook.save(path)

    def _make_entry(self, root: Path, report_path: Path) -> ChronologyEntry:
        bin_path = root / "2. Bin File" / "Master" / "Initial" / "FW_LRPS480_BL20_V1.00.bin"
        bin_path.parent.mkdir(parents=True, exist_ok=True)
        bin_path.write_bytes(b"demo")

        return ChronologyEntry(
            source_code_path=bin_path.resolve(),
            bin_file_path=bin_path,
            sdoc_file_path=None,
            bin_file_name=bin_path.name,
            sdoc_file_name=None,
            version="V1.00",
            plc_model="LRPS480",
            selpro_version="",
            bootloader_version="BL20",
            crc="ABCD1234",
            testing_stage="Master Initial",
            release_date="",
            reason_for_upgrade="",
            test_report_path=report_path,
        )

    def test_series_mode_writes_product_and_applicable_products_with_reload_proof(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            template_path = root / "template.xlsx"
            output_path = root / "Ladder_Chronology.xlsx"
            report_path = root / "TR-001-sgn.pdf"
            report_path.write_bytes(b"pdf")

            self._create_template(template_path)
            chronology = ProjectChronology(
                project_root=root,
                entries=[self._make_entry(root, report_path)],
            )

            series_info = ProductInfo(
                source_pdf=report_path,
                mode="Series",
                series_name="LRPS480",
                single_product=None,
                products=[
                    "LRPS480-24-CE",
                    "LRPS480-12-CE",
                    "LRPS480-05-CE",
                    "LRPS480-48-CE",
                ],
                page_number=1,
                parse_status="Success",
            )

            writer = ChronologyExcelWriter(template_path)
            with patch("services.chronology_generator.excel_writer.TestReportParser.extract_product_info", return_value=series_info):
                writer.write(chronology, output_path)

            # Reopen from disk: values must be physically persisted.
            wb = openpyxl.load_workbook(output_path)
            ws = wb.active

            self.assertEqual(ws["B3"].value, "LRPS480")
            self.assertEqual(ws["A9"].value, 1)
            self.assertEqual(ws["B9"].value, "LRPS480-24-CE")
            self.assertEqual(ws["A10"].value, 2)
            self.assertEqual(ws["B10"].value, "LRPS480-12-CE")
            self.assertEqual(ws["A11"].value, 3)
            self.assertEqual(ws["B11"].value, "LRPS480-05-CE")
            self.assertEqual(ws["A12"].value, 4)
            self.assertEqual(ws["B12"].value, "LRPS480-48-CE")

            # Header moved down because one row was inserted above the history table.
            self.assertEqual(ws["A13"].value, "Serial No.")
            wb.close()

    def test_single_mode_writes_product_value_with_reload_proof(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            template_path = root / "template.xlsx"
            output_path = root / "Ladder_Chronology.xlsx"
            report_path = root / "TR-002-sgn.pdf"
            report_path.write_bytes(b"pdf")

            self._create_template(template_path)
            chronology = ProjectChronology(
                project_root=root,
                entries=[self._make_entry(root, report_path)],
            )

            single_info = ProductInfo(
                source_pdf=report_path,
                mode="Single",
                series_name=None,
                single_product="LRPS480-24-CE",
                products=[],
                page_number=1,
                parse_status="Success",
            )

            writer = ChronologyExcelWriter(template_path)
            with patch("services.chronology_generator.excel_writer.TestReportParser.extract_product_info", return_value=single_info):
                writer.write(chronology, output_path)

            wb = openpyxl.load_workbook(output_path)
            ws = wb.active

            self.assertEqual(ws["B3"].value, "LRPS480-24-CE")
            # Single mode should not insert product rows above the history table.
            self.assertEqual(ws["A12"].value, "Serial No.")
            wb.close()


if __name__ == "__main__":
    unittest.main()
