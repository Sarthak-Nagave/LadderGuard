import shutil
from pathlib import Path

import openpyxl

from services.chronology_generator.generator import ChronologyGenerator


def create_template(path: Path, num_history: int, versions: list | None = None):
    wb = openpyxl.Workbook()
    ws = wb.active
    
    headers = {
        1: 'Serial No.', 2: 'Source Code Path', 3: 'Bin File Name', 4: 'CRC', 
        5: 'Ladder Version No.', 6: 'Reason For Upgrade', 7: 'Testing Stage', 
        8: 'PLC Model', 9: 'Selpro Version & Path', 10: 'Bootloader Version',
        11: 'Release Date', 12: 'Released By', 13: 'Ladder Release-To Production',
        14: 'Operator Procedure Modification', 15: 'Automation Set Up Modification', 16: 'Tested By'
    }
    for col, val in headers.items():
        ws.cell(row=14, column=col, value=val)
        
    next_row = 15
    for i in range(num_history):
        version = versions[i] if versions else f"V{num_history - i}.00"
        data = {
            1: num_history - i,
            2: f"C:\\Path\\{version}",
            3: f"FW_{version}.bin",
            4: "ABCD1234",
            5: version,
            6: "Release",
            7: "Production",
            8: "MODEL-X"
        }
        for col, val in data.items():
            ws.cell(row=next_row, column=col, value=val)
        next_row += 1
        
    ws.cell(row=next_row, column=2, value="Instructions")
    ws.merge_cells(f"B{next_row}:P{next_row}")
    
    ws.cell(row=next_row + 4, column=2, value="Footer Text")
    ws.merge_cells(f"B{next_row + 4}:D{next_row + 4}")
    
    wb.save(path)
    wb.close()

def setup_e2e_project(root: Path, template_path: Path):
    if root.exists():
        shutil.rmtree(root)
    root.mkdir(parents=True)
    
    bin_dir = root / "2. Bin File"
    chron_dir = root / "7. Chronology"
    
    for d in ["Master", "Slave", "UUT"]:
        (bin_dir / d).mkdir(parents=True)
        (chron_dir / d).mkdir(parents=True)
        
    # Master: 1 existing history (V1.01) + new V1.02
    create_template(chron_dir / "Master" / "Ladder_Chronology.xlsx", 1, ["V1.01"])
    (bin_dir / "Master" / "FW_V1.02.bin").touch()
    (bin_dir / "Master" / "FW_V1.02.sdoc").touch()

    # Slave: brand new chronology (needs base template)
    create_template(template_path, 0)
    (bin_dir / "Slave" / "FW_V2.01.bin").touch()
    (bin_dir / "Slave" / "FW_V2.01.sdoc").touch()
    
    # UUT: multiple history rows
    create_template(chron_dir / "UUT" / "Ladder_Chronology.xlsx", 2, ["V3.00", "V2.00"])
    (bin_dir / "UUT" / "FW_V3.01.bin").touch()
    (bin_dir / "UUT" / "FW_V3.01.sdoc").touch()

def run_e2e():
    root = Path("mock_e2e_project")
    base_template = Path("base_template.xlsx")
    setup_e2e_project(root, base_template)
    
    generator = ChronologyGenerator(root, base_template)
    
    print("=== SCENARIO 1: Generate all chronology files in one execution ===")
    generator.generate()
    
    print("\nMaster History:")
    print_sheet_state(root / "7. Chronology" / "Master" / "Ladder_Chronology.xlsx")
    
    print("\nSlave History:")
    print_sheet_state(root / "7. Chronology" / "Slave" / "Ladder_Chronology.xlsx")
    
    print("\nUUT History:")
    print_sheet_state(root / "7. Chronology" / "UUT" / "Ladder_Chronology.xlsx")
    
    print("\n=== SCENARIO 2: Run chronology generation twice ===")
    generator.generate()
    print("Master History after 2nd run:")
    print_sheet_state(root / "7. Chronology" / "Master" / "Ladder_Chronology.xlsx")
    
    print("\n=== SCENARIO 3: Run after another version release ===")
    (root / "2. Bin File" / "Master" / "FW_V1.03.bin").touch()
    (root / "2. Bin File" / "Master" / "FW_V1.03.sdoc").touch()
    generator.generate()
    print("Master History after V1.03:")
    print_sheet_state(root / "7. Chronology" / "Master" / "Ladder_Chronology.xlsx")
    
    print("\n=== SCENARIO 4: Generate 20 consecutive releases ===")
    for i in range(4, 24): # V1.04 to V1.23
        (root / "2. Bin File" / "Master" / f"FW_V1.{i:02d}.bin").touch()
        (root / "2. Bin File" / "Master" / f"FW_V1.{i:02d}.sdoc").touch()
        generator.generate()
    
    print("Master History after 20 releases:")
    print_sheet_state(root / "7. Chronology" / "Master" / "Ladder_Chronology.xlsx")
    
    print("\n=== SCENARIO 5: Verify Integrity (Open in Openpyxl) ===")
    try:
        wb = openpyxl.load_workbook(root / "7. Chronology" / "Master" / "Ladder_Chronology.xlsx")
        wb.close()
        print("Success: Workbook loaded without corruption.")
    except Exception as e:
        print(f"Error loading workbook: {e}")

def print_sheet_state(path: Path):
    if not path.exists():
        print(f"File not found: {path}")
        return
        
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    
    rows = []
    for r in range(15, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        # if there is a serial number or it's not None
        if v or ws.cell(row=r, column=2).value:
            rows.append([ws.cell(row=r, column=c).value for c in range(1, 6)])
            
    history_count = sum(1 for row in rows if row[0] is not None and isinstance(row[0], int))
    
    print(f"Total History Rows (with Serial): {history_count}")
    if len(rows) > 5:
        for row in rows[:3]: print(row)
        print("...")
        for row in rows[-2:]: print(row)
    else:
        for row in rows: print(row)
    wb.close()

if __name__ == "__main__":
    run_e2e()
