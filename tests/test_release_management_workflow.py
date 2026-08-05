from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
from PySide6.QtWidgets import QApplication

from gui.chronology.chronology_dialog import ChronologyDialog


class ReleaseManagementWorkflowTests(unittest.TestCase):
    def setUp(self) -> None:
        self.app = QApplication.instance() or QApplication([])

    def _create_template(self, path: Path) -> None:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        headers = [
            "Serial No.",
            "Source Code Path",
            "Bin File Name",
            "CRC",
            "Ladder Version No.",
            "Reason for Upgrade",
            "Testing Stage",
            "PLC Model",
            "Selpro Version & Path",
            "Bootloader Version",
            "Release Date",
            "Released By",
            "Ladder Release-To Production",
            "Operator Procedure Modification",
            "Automation Set Up Modification",
            "Tested By",
        ]
        for idx, header in enumerate(headers, start=1):
            sheet.cell(row=1, column=idx).value = header
        workbook.save(path)

    def _fill_required_fields(self, dialog: ChronologyDialog) -> None:
        for index, card in enumerate(dialog._entry_cards, start=1):
            card["line_reason"].setText(f"Release note {index}")
            card["line_released_by"].setText(f"Dev {index}")
            card["line_tested_by"].setText(f"QA {index}")
            card["combo_ladder_release"].setCurrentText("Yes")
            card["combo_operator_modification"].setCurrentText("No")
            card["combo_automation_modification"].setCurrentText("No")

    def test_all_pass_firmware_folders_generate_independent_chronologies(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            template_path = root / "template.xlsx"
            self._create_template(template_path)

            stage1 = root / "2. Bin File" / "Master" / "Initial"
            stage2 = root / "2. Bin File" / "Slave" / "Final"
            stage3 = root / "2. Bin File" / "UUT" / "QC"
            stage1.mkdir(parents=True, exist_ok=True)
            stage2.mkdir(parents=True, exist_ok=True)
            stage3.mkdir(parents=True, exist_ok=True)

            bin1 = stage1 / "FW_MST_INITIAL_MIBRX-4M_BL20_V1.00.bin"
            bin2 = stage2 / "FW_SLV_FINAL_FLEXYS-2M_BL10_V1.01.bin"
            bin3 = stage3 / "FW_UUT_QC_MIBRX-2M_BL30_V1.02.bin"
            bin1.write_bytes(b"1")
            bin2.write_bytes(b"2")
            bin3.write_bytes(b"3")

            records = {
                "Master/Initial": {"status": "PASS", "bin_file": str(bin1), "bin_name": bin1.name, "bin_crc": "AAAABBBB"},
                "Slave/Final": {"status": "PASS", "bin_file": str(bin2), "bin_name": bin2.name, "bin_crc": "CCCCDDDD"},
                "UUT/QC": {"status": "PASS", "bin_file": str(bin3), "bin_name": bin3.name, "bin_crc": "EEEEFFFF"},
            }

            dialog = ChronologyDialog(project_folder=root, validation_bin_crc_records=records)
            self._fill_required_fields(dialog)

            with patch.object(
                ChronologyDialog,
                "_selected_template_path_for_card",
                return_value=template_path,
            ), patch("gui.chronology.chronology_dialog.QMessageBox.information") as info_mock, patch(
                "gui.chronology.chronology_dialog.QMessageBox.warning"
            ) as warn_mock, patch(
                "gui.chronology.chronology_dialog.QMessageBox.critical"
            ) as crit_mock:
                dialog._on_generate_clicked()

            self.assertTrue((root / "7. Chronology" / "Master" / "Initial" / "Ladder_Chronology.xlsx").exists())
            self.assertTrue((root / "7. Chronology" / "Slave" / "Final" / "Ladder_Chronology.xlsx").exists())
            self.assertTrue((root / "7. Chronology" / "UUT" / "QC" / "Ladder_Chronology.xlsx").exists())

            title = info_mock.call_args[0][1]
            message = info_mock.call_args[0][2]
            self.assertEqual(title, "Chronology Generation Completed")
            self.assertIn("Generated: 3", message)
            self.assertIn("Skipped: 0", message)

    def test_partial_generation_skips_failed_firmware_folders(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            template_path = root / "template.xlsx"
            self._create_template(template_path)

            stage1 = root / "2. Bin File" / "Master" / "Initial"
            stage2 = root / "2. Bin File" / "Slave" / "Final"
            stage3 = root / "2. Bin File" / "UUT" / "QC"
            stage1.mkdir(parents=True, exist_ok=True)
            stage2.mkdir(parents=True, exist_ok=True)
            stage3.mkdir(parents=True, exist_ok=True)

            bin1 = stage1 / "FW_MST_INITIAL_MIBRX-4M_BL20_V1.00.bin"
            bin2 = stage2 / "FW_SLV_FINAL_FLEXYS-2M_BL10_V1.01.bin"
            bin3 = stage3 / "FW_UUT_QC_MIBRX-2M_BL30_V1.02.bin"
            bin1.write_bytes(b"1")
            bin2.write_bytes(b"2")
            bin3.write_bytes(b"3")

            records = {
                "Master/Initial": {"status": "PASS", "bin_file": str(bin1), "bin_name": bin1.name, "bin_crc": "AAAABBBB"},
                "Slave/Final": {"status": "PASS", "bin_file": str(bin2), "bin_name": bin2.name, "bin_crc": "CCCCDDDD"},
                "UUT/QC": {"status": "FAILED", "bin_file": str(bin3), "bin_name": bin3.name, "bin_crc": "EEEEFFFF", "reason": "CRC Mismatch"},
            }

            dialog = ChronologyDialog(project_folder=root, validation_bin_crc_records=records)
            self._fill_required_fields(dialog)

            with patch.object(
                ChronologyDialog,
                "_selected_template_path_for_card",
                return_value=template_path,
            ), patch("gui.chronology.chronology_dialog.QMessageBox.information") as info_mock, patch(
                "gui.chronology.chronology_dialog.QMessageBox.warning"
            ) as warn_mock, patch(
                "gui.chronology.chronology_dialog.QMessageBox.critical"
            ) as crit_mock:
                dialog._on_generate_clicked()

            self.assertTrue((root / "7. Chronology" / "Master" / "Initial" / "Ladder_Chronology.xlsx").exists())
            self.assertTrue((root / "7. Chronology" / "Slave" / "Final" / "Ladder_Chronology.xlsx").exists())
            self.assertFalse((root / "7. Chronology" / "UUT" / "QC" / "Ladder_Chronology.xlsx").exists())

            message = info_mock.call_args[0][2]
            self.assertIn("Generated: 2", message)
            self.assertIn("Skipped: 1", message)
            self.assertIn("UUT/QC", message)
            self.assertIn("CRC Mismatch", message)

    def test_existing_chronology_workbook_preserves_history_when_updated(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            template_path = root / "template.xlsx"
            self._create_template(template_path)

            stage = root / "2. Bin File" / "Master" / "Initial"
            stage.mkdir(parents=True, exist_ok=True)
            selected_bin = stage / "FW_MST_INITIAL_MIBRX-4M_BL20_V1.01.bin"
            selected_bin.write_bytes(b"new")

            chronology_path = root / "7. Chronology" / "Master" / "Initial" / "Ladder_Chronology.xlsx"
            chronology_path.parent.mkdir(parents=True, exist_ok=True)
            self._create_template(chronology_path)

            workbook = openpyxl.load_workbook(chronology_path)
            sheet = workbook.active
            sheet.cell(row=2, column=1).value = 1
            sheet.cell(row=2, column=2).value = str((stage / "FW_MST_INITIAL_MIBRX-4M_BL20_V1.00.bin").resolve())
            sheet.cell(row=2, column=3).value = "FW_MST_INITIAL_MIBRX-4M_BL20_V1.00.bin"
            sheet.cell(row=2, column=4).value = "AABBCCDD"
            sheet.cell(row=2, column=5).value = "V1.00"
            sheet.cell(row=2, column=6).value = "Previous release"
            sheet.cell(row=2, column=7).value = "Master Initial"
            sheet.cell(row=2, column=8).value = "MIBRX-4M"
            sheet.cell(row=2, column=10).value = "BL20"
            sheet.cell(row=2, column=11).value = "01/07/2026"
            sheet.cell(row=2, column=3).font = openpyxl.styles.Font(bold=True)
            sheet.cell(row=2, column=3).fill = openpyxl.styles.PatternFill("solid", fgColor="D9EAF7")
            sheet.cell(row=2, column=3).border = openpyxl.styles.Border(
                left=openpyxl.styles.Side(style="thin"),
                right=openpyxl.styles.Side(style="thin"),
                top=openpyxl.styles.Side(style="thin"),
                bottom=openpyxl.styles.Side(style="thin"),
            )
            sheet.cell(row=2, column=3).alignment = openpyxl.styles.Alignment(horizontal="center")
            sheet.row_dimensions[2].height = 24
            workbook.save(chronology_path)

            records = {
                "Master/Initial": {"status": "PASS", "bin_file": str(selected_bin), "bin_name": selected_bin.name, "bin_crc": "CCCCDDDD"},
            }

            dialog = ChronologyDialog(project_folder=root, validation_bin_crc_records=records)
            self._fill_required_fields(dialog)

            with patch("gui.chronology.chronology_dialog.QMessageBox.information"), patch(
                "gui.chronology.chronology_dialog.QMessageBox.warning"
            ), patch("gui.chronology.chronology_dialog.QMessageBox.critical"):
                dialog._on_generate_clicked()

            workbook = openpyxl.load_workbook(chronology_path)
            sheet = workbook.active
            self.assertEqual(sheet.cell(row=2, column=1).value, 2)
            self.assertEqual(sheet.cell(row=2, column=3).value, selected_bin.name)
            self.assertEqual(sheet.cell(row=3, column=1).value, 1)
            self.assertEqual(sheet.cell(row=3, column=3).value, "FW_MST_INITIAL_MIBRX-4M_BL20_V1.00.bin")
            self.assertEqual(sheet.cell(row=2, column=4).value, "CCCCDDDD")
            self.assertEqual(sheet.cell(row=3, column=4).value, "AABBCCDD")
            self.assertEqual(sheet.row_dimensions[2].height, 24)
            self.assertTrue(sheet.cell(row=2, column=3).font.bold)
            self.assertEqual(sheet.cell(row=2, column=3).fill.fgColor.rgb, "00D9EAF7")
            self.assertEqual(sheet.cell(row=2, column=3).border.left.style, "thin")
            self.assertEqual(sheet.cell(row=2, column=3).alignment.horizontal, "center")


if __name__ == "__main__":
    unittest.main()
